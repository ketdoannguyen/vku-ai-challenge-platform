"""Hai nhánh Public/Private: hợp đồng tạo/đọc, lịch nhánh, GT theo nhánh, khóa và công bố.

Nhóm test này khóa ranh giới của các giai đoạn A-D: domain contract, cấu hình nhánh, GT riêng,
bằng chứng xác minh theo nhánh, khóa scoring, quyền nộp/quota/vòng đời lượt, và riêng tư kết quả
Private chưa công bố trên mọi bề mặt thí sinh.
"""

import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from bson import ObjectId
from fastapi import HTTPException
from openpyxl import load_workbook

from app.accounts.service import ACCOUNTS_COLLECTION
from app.competitions.service import COMPETITIONS_COLLECTION
from app.core.config import get_settings
from app.ai_review import content_snapshot
from app.core.datetimes import utc_day_key
from app.memberships.service import MEMBERSHIPS_COLLECTION
from app.scoring_attempts import service as attempts_service
from app.scoring_attempts import store as attempts_store
from app.scoring_attempts import worker
from app.scoring_attempts.store import ATTEMPTS_COLLECTION, STATUS_QUEUED, STATUS_STAGING
from app.submissions import service as submissions_service
from tests.ai_review_helpers import ai_env  # noqa: F401 - fixture cho test AI theo nhánh
from tests.helpers import (
    GROUND_TRUTH_CSV,
    PARTICIPANT_CREDENTIALS,
    SCORING_CONFIG,
    SUBMISSION_GROUND_TRUTH,
    V2_CONTRACT,
    V2_GROUND_TRUTH,
    V2_SUBMISSION,
    VALID_NOTEBOOK,
    attempt_documents,
    attempt_status,
    login,
    login_participant,
    membership_document,
    publish_v2_competition,
    put_scoring_v2,
    ready_competition,
    run_worker,
    submission_documents,
    submit,
)
from tests.test_ai_review_api import _add_rules, enable_ai


def at(value: str) -> datetime:
    """Parse timestamp API (hậu tố Z) về datetime aware để so sánh."""
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def same_time(actual: str, expected: datetime) -> bool:
    """So timestamp API với mốc kỳ vọng; bỏ microsecond vì Mongo chỉ lưu tới mili giây."""
    return at(actual).replace(microsecond=0) == expected.replace(microsecond=0)


def iso(value: datetime) -> str:
    return value.isoformat()


def err(response) -> dict:
    """Payload lỗi theo format chung của API: code/message/detail nằm trong `error`."""
    return response.json()["error"]


def near(actual: timedelta, expected: timedelta) -> bool:
    return abs(actual - expected) < timedelta(seconds=1)


def flipped_v2_ground_truth() -> bytes:
    """Bộ GT thứ hai hợp lệ nhưng khác nội dung - dùng để chứng minh bằng chứng stale theo nhánh."""
    return b"id,label,predict_type\n1,1,B\n2,0,A\n3,1,B\n4,0,A\n"


def dual_body(
    slug: str = "dual-cup",
    *,
    public_hours: tuple[int, int] = (-1, 1),
    private_hours: tuple[int, int] = (-1, 1),
    quota_per_day: int = 5,
    result_policy: str | None = None,
    publish_condition: str | None = None,
    private_resources: list[dict] | None = None,
) -> dict:
    now = datetime.now(timezone.utc)
    public = {
        "start_at": iso(now + timedelta(hours=public_hours[0])),
        "end_at": iso(now + timedelta(hours=public_hours[1])),
        "quota_per_day": quota_per_day,
    }
    private = {
        "start_at": iso(now + timedelta(hours=private_hours[0])),
        "end_at": iso(now + timedelta(hours=private_hours[1])),
        "quota_per_day": quota_per_day,
        "resources": private_resources or [],
    }
    if result_policy is not None:
        private["result_policy"] = result_policy
    if publish_condition is not None:
        private["publish_condition"] = publish_condition
    return {
        "slug": slug,
        "name": slug,
        "mode": "public_private",
        "primary_metric": "f1",
        "public_track": public,
        "private_track": private,
    }


def create_dual(client, body: dict):
    return client.post("/api/admin/competitions", json=body)


def create_dual_ok(client, body: dict | None = None) -> dict:
    response = create_dual(client, body or dual_body())
    assert response.status_code == 201, response.text
    return response.json()


def upload_gt(
    client,
    competition_id: str,
    data: bytes,
    *,
    track: str | None = None,
    expected_revision: int | None = None,
):
    form = {}
    if track is not None:
        form["track"] = track
    if expected_revision is not None:
        form["expected_revision"] = str(expected_revision)
    return client.put(
        f"/api/admin/competitions/{competition_id}/ground-truth",
        files={"file": ("ground_truth.csv", data, "text/csv")},
        data=form,
    )


def run_test(
    client,
    competition_id: str,
    *,
    track: str | None = None,
    data: bytes = V2_SUBMISSION,
    expected_revision: int | None = None,
):
    form = {}
    if track is not None:
        form["track"] = track
    if expected_revision is not None:
        form["expected_revision"] = str(expected_revision)
    return client.post(
        f"/api/admin/competitions/{competition_id}/scoring/test",
        files={"file": ("sample.csv", data, "text/csv")},
        data=form,
    )


def publish_dual_v1(
    client,
    competition_id: str,
    *,
    slug: str,
    public_gt: bytes = GROUND_TRUTH_CSV,
    private_gt: bytes = SUBMISSION_GROUND_TRUTH,
    join: bool = True,
) -> dict:
    """Đưa cuộc thi dual qua cổng publish bằng bộ chấm v1 với hai bộ GT riêng; trả detail sau publish.

    `join=False` cho các cuộc thi đã hết hạn tham gia ngay lúc tạo - test thuần admin không cần
    thành viên và join sẽ bị luật deadline chặn.
    """
    configured = client.put(
        f"/api/admin/competitions/{competition_id}/scoring", json=SCORING_CONFIG
    )
    assert configured.status_code == 200, configured.text
    assert upload_gt(client, competition_id, public_gt, track="public").status_code == 200
    assert upload_gt(client, competition_id, private_gt, track="private").status_code == 200
    published = client.post(f"/api/admin/competitions/{competition_id}/publish")
    assert published.status_code == 200, published.text
    if join:
        login_participant(client)
        joined = client.post(f"/api/competitions/{slug}/join", json={})
        assert joined.status_code == 200, joined.text
        login(client)
    return published.json()


# --- Tạo và validate ------------------------------------------------------------------------------


def test_create_dual_requires_both_tracks(client):
    login(client)
    body = dual_body()
    del body["private_track"]
    response = create_dual(client, body)
    assert response.status_code == 422
    assert "Public" in err(response)["message"] and "Private" in err(response)["message"]

    body = dual_body()
    del body["public_track"]
    assert create_dual(client, body).status_code == 422


def test_create_dual_rejects_top_level_schedule(client):
    login(client)
    now = datetime.now(timezone.utc)
    body = dual_body()
    body["start_at"] = iso(now - timedelta(hours=1))
    body["end_at"] = iso(now + timedelta(hours=1))
    response = create_dual(client, body)
    assert response.status_code == 422
    assert "lịch của từng nhánh" in err(response)["message"]


def test_create_single_rejects_track_config(client):
    login(client)
    now = datetime.now(timezone.utc)
    response = create_dual(
        client,
        {
            "slug": "single-with-tracks",
            "name": "single",
            "start_at": iso(now - timedelta(hours=1)),
            "end_at": iso(now + timedelta(hours=1)),
            "public_track": {"start_at": iso(now), "end_at": iso(now + timedelta(hours=1))},
        },
    )
    assert response.status_code == 422
    assert "thông thường" in err(response)["message"]


def test_create_rejects_unknown_mode(client):
    login(client)
    now = datetime.now(timezone.utc)
    response = create_dual(
        client,
        {
            "slug": "bad-mode",
            "name": "bad",
            "mode": "two-way",
            "start_at": iso(now - timedelta(hours=1)),
            "end_at": iso(now + timedelta(hours=1)),
        },
    )
    assert response.status_code == 422


def test_create_dual_validates_track_schedule_and_quota(client):
    login(client)
    body = dual_body()
    body["public_track"]["end_at"] = body["public_track"]["start_at"]
    assert create_dual(client, body).status_code == 422

    body = dual_body()
    body["private_track"]["quota_per_day"] = 1001
    response = create_dual(client, body)
    assert response.status_code == 422
    assert "Private" in err(response)["message"]

    assert create_dual(client, dual_body(result_policy="sometimes")).status_code == 422
    assert create_dual(client, dual_body(publish_condition="maybe")).status_code == 422


def test_create_dual_document_shape(client):
    login(client)
    now = datetime.now(timezone.utc)
    payload = create_dual_ok(
        client,
        dual_body(slug="shape-cup", public_hours=(-2, 1), private_hours=(-1, 3), quota_per_day=7),
    )
    assert payload["mode"] == "public_private"
    # Envelope: sớm nhất trong các giờ mở, muộn nhất trong các giờ đóng.
    assert same_time(payload["start_at"], now + timedelta(hours=-2))
    assert same_time(payload["end_at"], now + timedelta(hours=3))
    # Quota cấp cuộc thi không tồn tại ở dual.
    assert payload["quota_per_day"] is None
    assert payload["control_revision"] == 1
    assert payload["stop_generation"] == 0
    assert payload["scoring_locked"] is False
    tracks = payload["tracks"]
    assert set(tracks) == {"public", "private"}
    assert tracks["public"]["quota_per_day"] == 7
    assert tracks["public"]["window_state"] == "open"
    assert tracks["private"]["result_policy"] == "manual"
    assert tracks["private"]["publish_condition"] == "admin_decides"
    assert tracks["private"]["results_released"] is False
    assert tracks["private"]["results_published_at"] is None
    assert tracks["public"]["admission_seq"] == 0
    assert tracks["private"]["admission_seq"] == 0
    assert tracks["public"]["ground_truth"] is None
    assert tracks["public"]["ready"] is False
    assert payload["publish_ready"] is False


def test_single_payload_keeps_legacy_shape(client):
    login(client)
    now = datetime.now(timezone.utc)
    created = client.post(
        "/api/admin/competitions",
        json={
            "slug": "plain-cup",
            "name": "plain",
            "start_at": iso(now - timedelta(hours=1)),
            "end_at": iso(now + timedelta(hours=1)),
            "primary_metric": "f1",
            "quota_per_day": 5,
        },
    )
    assert created.status_code == 201
    payload = created.json()
    assert payload["mode"] == "single"
    assert payload["tracks"] is None
    assert payload["quota_per_day"] == 5
    # Single không có điều khiển dual nào trong payload.
    assert "control_revision" not in payload
    assert "stop_generation" not in payload


# --- Đọc công khai và tài nguyên nhánh ------------------------------------------------------------


def test_public_landing_exposes_windows_and_gates_track_resources(client):
    login(client)
    resource = {"label": "Test Private", "url": "https://drive.google.com/file/d/private-test"}
    payload = create_dual_ok(
        client,
        dual_body(
            slug="landing-cup",
            public_hours=(2, 4),  # chưa mở
            private_hours=(-1, 4),  # đang mở
            private_resources=[resource],
        ),
    )
    publish_dual_v1(client, payload["id"], slug="landing-cup")

    # Khách chưa đăng nhập: thấy lịch và cửa sổ, không có URL tài nguyên nào.
    client.cookies.clear()
    public = client.get("/api/competitions/landing-cup")
    assert public.status_code == 200
    landing = public.json()
    assert landing["mode"] == "public_private"
    assert set(landing["tracks"]) == {"public", "private"}
    assert landing["tracks"]["public"]["window_state"] == "scheduled"
    assert landing["tracks"]["private"]["window_state"] == "open"
    assert landing["tracks"]["private"]["results_released"] is False
    assert "resources" not in landing["tracks"]["private"]

    # Thành viên: tài nguyên nhánh chỉ rời backend từ giờ mở của chính nhánh đó.
    login_participant(client)
    member = client.get("/api/competitions/landing-cup").json()
    assert member["tracks"]["public"]["resources"] == []
    assert member["tracks"]["private"]["resources"] == [resource]
    assert member["tracks"]["public"]["can_submit"] is False
    assert member["tracks"]["public"]["blocked_reason"] == "not_open"
    assert member["tracks"]["private"]["can_submit"] is True
    assert member["tracks"]["private"]["submission_ready"] is True


def test_participant_track_windows_use_server_clock(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="closed-public", public_hours=(-3, -1)))
    publish_dual_v1(client, payload["id"], slug="closed-public")
    login_participant(client)
    detail = client.get("/api/competitions/closed-public").json()
    assert detail["tracks"]["public"]["window_state"] == "closed"
    assert detail["tracks"]["public"]["can_submit"] is False
    assert detail["tracks"]["public"]["blocked_reason"] == "deadline_passed"


# --- Sửa cấu hình và lịch nhánh --------------------------------------------------------------------


def test_dual_edit_blocks_top_level_schedule_and_quota(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="edit-cup"))
    cid = payload["id"]
    name_only = client.patch(f"/api/admin/competitions/{cid}", json={"name": "Tên mới"})
    assert name_only.status_code == 200, name_only.text
    assert name_only.json()["name"] == "Tên mới"

    now = datetime.now(timezone.utc)
    for body in (
        {"start_at": iso(now)},
        {"end_at": iso(now + timedelta(days=2))},
        {"quota_per_day": 3},
    ):
        response = client.patch(f"/api/admin/competitions/{cid}", json=body)
        assert response.status_code == 422
        assert "lịch nhánh" in err(response)["message"]


def test_track_schedule_patch_updates_envelope_and_revision(client):
    login(client)
    payload = create_dual_ok(
        client, dual_body(slug="schedule-cup", public_hours=(-2, 1), private_hours=(-1, 2))
    )
    cid = payload["id"]
    now = datetime.now(timezone.utc)
    extended_end = now + timedelta(hours=6)

    response = client.patch(
        f"/api/admin/competitions/{cid}/tracks/public/schedule",
        json={
            "start_at": iso(now - timedelta(hours=2)),
            "end_at": iso(extended_end),
            "expected_revision": 1,
            "reason": "Gia hạn nhánh Public theo nguyện vọng đội thi",
        },
    )
    assert response.status_code == 200, response.text
    updated = response.json()
    assert updated["control_revision"] == 2
    assert same_time(updated["tracks"]["public"]["end_at"], extended_end)
    assert updated["last_change"]["action"] == "track_schedule"
    assert updated["last_change"]["revision"] == 2
    assert updated["last_change"]["reason"] == "Gia hạn nhánh Public theo nguyện vọng đội thi"

    # Lịch cũ (revision cũ) không ghi đè được nữa.
    stale = client.patch(
        f"/api/admin/competitions/{cid}/tracks/public/schedule",
        json={
            "start_at": iso(now),
            "end_at": iso(now + timedelta(hours=1)),
            "expected_revision": 1,
            "reason": "ghi đè",
        },
    )
    assert stale.status_code == 409
    assert err(stale)["detail"]["control_revision"] == 2

    # Thiếu lý do hoặc sai nhánh không đi được.
    missing_reason = client.patch(
        f"/api/admin/competitions/{cid}/tracks/private/schedule",
        json={
            "start_at": iso(now),
            "end_at": iso(now + timedelta(hours=2)),
            "expected_revision": 2,
            "reason": "   ",
        },
    )
    assert missing_reason.status_code == 422
    bad_track = client.patch(
        f"/api/admin/competitions/{cid}/tracks/third/schedule",
        json={
            "start_at": iso(now),
            "end_at": iso(now + timedelta(hours=2)),
            "expected_revision": 2,
            "reason": "sai nhánh",
        },
    )
    assert bad_track.status_code == 422
    assert err(bad_track)["code"] == "INVALID_TRACK"


def test_track_schedule_patch_rejects_single_competition(client):
    login(client)
    now = datetime.now(timezone.utc)
    created = client.post(
        "/api/admin/competitions",
        json={
            "slug": "single-schedule",
            "name": "single",
            "start_at": iso(now - timedelta(hours=1)),
            "end_at": iso(now + timedelta(hours=1)),
            "primary_metric": "f1",
            "quota_per_day": 5,
        },
    )
    assert created.status_code == 201
    response = client.patch(
        f"/api/admin/competitions/{created.json()['id']}/tracks/public/schedule",
        json={
            "start_at": iso(now),
            "end_at": iso(now + timedelta(hours=2)),
            "expected_revision": 1,
            "reason": "không có nhánh",
        },
    )
    assert response.status_code == 422
    assert err(response)["code"] == "INVALID_TRACK"


# --- GT theo nhánh và cổng publish ----------------------------------------------------------------


def test_ground_truth_track_is_required_and_unknown_track_rejected(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="gt-track-cup"))
    cid = payload["id"]
    client.put(f"/api/admin/competitions/{cid}/scoring", json=SCORING_CONFIG)

    missing = upload_gt(client, cid, GROUND_TRUTH_CSV)
    assert missing.status_code == 422
    assert err(missing)["code"] == "TRACK_REQUIRED"

    unknown = upload_gt(client, cid, GROUND_TRUTH_CSV, track="secret")
    assert unknown.status_code == 422
    assert err(unknown)["code"] == "INVALID_TRACK"

    now = datetime.now(timezone.utc)
    single = client.post(
        "/api/admin/competitions",
        json={
            "slug": "gt-single",
            "name": "single",
            "start_at": iso(now - timedelta(hours=1)),
            "end_at": iso(now + timedelta(hours=1)),
            "primary_metric": "f1",
            "quota_per_day": 5,
        },
    ).json()
    client.put(f"/api/admin/competitions/{single['id']}/scoring", json=SCORING_CONFIG)
    with_track = upload_gt(client, single["id"], GROUND_TRUTH_CSV, track="public")
    assert with_track.status_code == 422
    assert err(with_track)["code"] == "INVALID_TRACK"


def test_dual_publish_requires_both_ground_truths(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="both-gt-cup"))
    cid = payload["id"]
    client.put(f"/api/admin/competitions/{cid}/scoring", json=SCORING_CONFIG)
    assert upload_gt(client, cid, GROUND_TRUTH_CSV, track="public").status_code == 200

    blocked = client.post(f"/api/admin/competitions/{cid}/publish")
    assert blocked.status_code == 422
    assert err(blocked)["code"] == "GROUND_TRUTH_REQUIRED"
    assert "Private" in err(blocked)["message"]

    assert upload_gt(client, cid, SUBMISSION_GROUND_TRUTH, track="private").status_code == 200
    published = client.post(f"/api/admin/competitions/{cid}/publish")
    assert published.status_code == 200, published.text
    assert published.json()["status"] == "published"


def test_public_test_does_not_verify_private_track(client, fake_runner):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="verify-cup"))
    cid = payload["id"]
    assert (
        put_scoring_v2(client, cid, expected_revision=0, output_contract=V2_CONTRACT).status_code
        == 200
    )
    missing_track = upload_gt(client, cid, V2_GROUND_TRUTH, expected_revision=1)
    assert missing_track.status_code == 422
    assert err(missing_track)["code"] == "TRACK_REQUIRED"
    assert (
        upload_gt(client, cid, V2_GROUND_TRUTH, track="public", expected_revision=1).status_code
        == 200
    )
    assert (
        upload_gt(
            client, cid, flipped_v2_ground_truth(), track="private", expected_revision=2
        ).status_code
        == 200
    )

    # Chỉ Public được chạy thử: Private vẫn thiếu bằng chứng.
    tested = run_test(client, cid, track="public", expected_revision=3)
    assert tested.status_code == 200, tested.text
    view = tested.json()
    assert view["tracks"]["public"]["verified"] is True
    assert view["tracks"]["private"]["verified"] is False
    assert view["tracks"]["private"]["not_ready_reason"]["code"] == "SCORING_TEST_REQUIRED"

    blocked = client.post(f"/api/admin/competitions/{cid}/publish")
    assert blocked.status_code == 422
    assert err(blocked)["code"] == "SCORING_TEST_REQUIRED"
    assert "Private" in err(blocked)["message"]

    assert run_test(client, cid, track="private", expected_revision=3).status_code == 200
    published = client.post(f"/api/admin/competitions/{cid}/publish")
    assert published.status_code == 200, published.text


def test_ground_truth_change_stales_only_that_track(client, fake_runner):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="stale-cup"))
    cid = payload["id"]
    assert (
        put_scoring_v2(client, cid, expected_revision=0, output_contract=V2_CONTRACT).status_code
        == 200
    )
    assert (
        upload_gt(client, cid, V2_GROUND_TRUTH, track="public", expected_revision=1).status_code
        == 200
    )
    assert (
        upload_gt(
            client, cid, flipped_v2_ground_truth(), track="private", expected_revision=2
        ).status_code
        == 200
    )
    for track in ("public", "private"):
        assert (
            run_test(
                client, cid, track=track, data=V2_SUBMISSION, expected_revision=3
            ).status_code
            == 200
        )
    view = client.get(f"/api/admin/competitions/{cid}/scoring").json()
    assert view["tracks"]["public"]["verified"] is True
    assert view["tracks"]["private"]["verified"] is True

    # Thay GT của Public làm bằng chứng của Public hết hiệu lực; Private còn nguyên.
    replaced = upload_gt(
        client, cid, flipped_v2_ground_truth(), track="public", expected_revision=3
    )
    assert replaced.status_code == 200, replaced.text
    view = replaced.json()
    assert view["tracks"]["public"]["verified"] is False
    assert view["tracks"]["private"]["verified"] is True
    assert view["tracks"]["private"]["verification"]["tested_by"] == "admin@vku.vn"


def test_scoring_locked_after_publish_and_reopen_does_not_unlock(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="lock-cup"))
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="lock-cup")
    assert client.get(f"/api/admin/competitions/{cid}").json()["scoring_locked"] is True

    locked_gt = upload_gt(client, cid, GROUND_TRUTH_CSV, track="public")
    assert locked_gt.status_code == 422
    assert err(locked_gt)["code"] == "SCORING_LOCKED"
    locked_config = client.put(f"/api/admin/competitions/{cid}/scoring", json=SCORING_CONFIG)
    assert locked_config.status_code == 422
    assert err(locked_config)["code"] == "SCORING_LOCKED"

    assert client.post(f"/api/admin/competitions/{cid}/close").status_code == 200
    assert client.post(f"/api/admin/competitions/{cid}/reopen").status_code == 200
    detail = client.get(f"/api/admin/competitions/{cid}").json()
    assert detail["scoring_locked"] is True
    assert detail["stop_generation"] == 1  # close là lệnh dừng quản trị
    assert detail["status"] == "published"


def test_reopen_with_track_schedules_updates_windows_atomically(client):
    login(client)
    payload = create_dual_ok(
        client, dual_body(slug="reopen-cup", public_hours=(-3, -1), private_hours=(-3, -1))
    )
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="reopen-cup", join=False)
    assert client.post(f"/api/admin/competitions/{cid}/close").status_code == 200
    revision = client.get(f"/api/admin/competitions/{cid}").json()["control_revision"]

    now = datetime.now(timezone.utc)
    private_end = now + timedelta(hours=4)
    no_reason = client.post(
        f"/api/admin/competitions/{cid}/reopen",
        json={
            "expected_revision": revision,
            "tracks": {"private": {"start_at": iso(now), "end_at": iso(private_end)}},
        },
    )
    assert no_reason.status_code == 422

    stale = client.post(
        f"/api/admin/competitions/{cid}/reopen",
        json={
            "expected_revision": revision - 1,
            "reason": "mở lại sai revision",
            "tracks": {"private": {"start_at": iso(now), "end_at": iso(private_end)}},
        },
    )
    assert stale.status_code == 409

    reopened = client.post(
        f"/api/admin/competitions/{cid}/reopen",
        json={
            "expected_revision": revision,
            "reason": "Mở lại để nhận bài bổ sung",
            "tracks": {
                "private": {"start_at": iso(now), "end_at": iso(private_end)},
                "public": {"start_at": iso(now), "end_at": iso(now + timedelta(hours=2))},
            },
        },
    )
    assert reopened.status_code == 200, reopened.text
    detail = reopened.json()
    assert detail["status"] == "published"
    assert same_time(detail["tracks"]["private"]["end_at"], private_end)
    assert same_time(detail["end_at"], private_end)  # envelope cập nhật cùng lượt ghi
    assert detail["tracks"]["private"]["window_state"] == "open"
    assert detail["control_revision"] == revision + 1


# --- Công bố kết quả Private ----------------------------------------------------------------------


def test_publish_results_manual_admin_decides_is_idempotent(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="release-cup"))
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="release-cup")
    revision = client.get(f"/api/admin/competitions/{cid}").json()["control_revision"]

    released = client.post(
        f"/api/admin/competitions/{cid}/tracks/private/publish-results",
        json={"expected_revision": revision, "reason": "Công bố kết quả chính thức"},
    )
    assert released.status_code == 200, released.text
    first_marker = released.json()["tracks"]["private"]["results_published_at"]
    assert first_marker is not None
    assert released.json()["tracks"]["private"]["results_published_by"] == "admin@vku.vn"

    # Retry sau khi đã công bố: thành công với đúng dấu mốc cũ, không cần revision khớp.
    retry = client.post(
        f"/api/admin/competitions/{cid}/tracks/private/publish-results",
        json={"expected_revision": 1, "reason": "double click"},
    )
    assert retry.status_code == 200
    assert retry.json()["tracks"]["private"]["results_published_at"] == first_marker


def test_released_results_cannot_go_back_to_hidden(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="hidden-cup", result_policy="immediate"))
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="hidden-cup")
    detail = client.get(f"/api/admin/competitions/{cid}").json()
    assert detail["tracks"]["private"]["results_released"] is True

    hide = client.patch(
        f"/api/admin/competitions/{cid}/tracks/private/policy",
        json={
            "result_policy": "manual",
            "expected_revision": detail["control_revision"],
            "reason": "thử giấu lại",
        },
    )
    assert hide.status_code == 409
    assert err(hide)["code"] == "RESULTS_ALREADY_RELEASED"


def test_private_release_monotonic_across_extend_and_reopen(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="monotonic-cup"))
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="monotonic-cup")
    revision = client.get(f"/api/admin/competitions/{cid}").json()["control_revision"]
    released = client.post(
        f"/api/admin/competitions/{cid}/tracks/private/publish-results",
        json={"expected_revision": revision},
    )
    assert released.status_code == 200, released.text
    marker = released.json()["tracks"]["private"]["results_published_at"]

    now = datetime.now(timezone.utc)
    extended = client.patch(
        f"/api/admin/competitions/{cid}/tracks/private/schedule",
        json={
            "start_at": iso(now),
            "end_at": iso(now + timedelta(hours=5)),
            "expected_revision": released.json()["control_revision"],
            "reason": "Gia hạn nhánh Private sau công bố",
        },
    )
    assert extended.status_code == 200, extended.text
    assert extended.json()["tracks"]["private"]["results_released"] is True

    assert client.post(f"/api/admin/competitions/{cid}/close").status_code == 200
    detail = client.get(f"/api/admin/competitions/{cid}").json()
    reopened = client.post(
        f"/api/admin/competitions/{cid}/reopen",
        json={
            "expected_revision": detail["control_revision"],
            "reason": "Mở lại giữ nguyên kết quả đã công bố",
            "tracks": {
                "private": {"start_at": iso(now), "end_at": iso(now + timedelta(hours=8))}
            },
        },
    )
    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["tracks"]["private"]["results_released"] is True
    assert reopened.json()["tracks"]["private"]["results_published_at"] == marker
    # Retry công bố sau reopen vẫn thành công và giữ dấu mốc đầu tiên.
    retry = client.post(
        f"/api/admin/competitions/{cid}/tracks/private/publish-results",
        json={"expected_revision": 1},
    )
    assert retry.status_code == 200
    assert retry.json()["tracks"]["private"]["results_published_at"] == marker


def test_strict_release_blocks_while_open(client):
    login(client)
    payload = create_dual_ok(
        client, dual_body(slug="strict-open-cup", publish_condition="after_closed_and_scored")
    )
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="strict-open-cup")
    revision = client.get(f"/api/admin/competitions/{cid}").json()["control_revision"]
    blocked = client.post(
        f"/api/admin/competitions/{cid}/tracks/private/publish-results",
        json={"expected_revision": revision},
    )
    assert blocked.status_code == 409
    assert err(blocked)["code"] == "RELEASE_CONDITION_NOT_MET"
    assert err(blocked)["detail"]["reason"] == "open"


def test_strict_release_blocks_while_processing(client):
    login(client)
    payload = create_dual_ok(
        client,
        dual_body(
            slug="strict-processing-cup",
            private_hours=(-3, -1),  # cửa Private đã đóng
            publish_condition="after_closed_and_scored",
        ),
    )
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="strict-processing-cup")
    revision = client.get(f"/api/admin/competitions/{cid}").json()["control_revision"]

    async def write_attempt(status: str) -> None:
        db = client.app.state.mongo.db
        await db[ATTEMPTS_COLLECTION].update_one(
            {"competition_id": ObjectId(cid), "track": "private"},
            {"$set": {"status": status}},
            upsert=True,
        )

    asyncio.run(write_attempt(STATUS_STAGING))
    blocked = client.post(
        f"/api/admin/competitions/{cid}/tracks/private/publish-results",
        json={"expected_revision": revision},
    )
    assert blocked.status_code == 409
    assert err(blocked)["detail"]["reason"] == "processing"
    assert err(blocked)["detail"]["active_attempts"] == 1

    asyncio.run(write_attempt("FAILED"))
    released = client.post(
        f"/api/admin/competitions/{cid}/tracks/private/publish-results",
        json={"expected_revision": revision},
    )
    assert released.status_code == 200, released.text


def test_scope_stats_dem_luot_theo_nhanh_va_trang_thai(client):
    """Số liệu trang công bố: completed/failed theo nhánh, in_flight đúng tập cổng chặt kiểm."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="stats-cup"))
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="stats-cup")
    # Public một bài hoàn tất; Private thêm một lượt hỏng và một lượt đang chờ.
    login_participant(client)
    assert submit(client, cid, V1_SUBMISSION, track="public").status_code == 201
    assert submit(client, cid, V1_SUBMISSION, track="private").status_code == 201

    async def write_attempt(track: str, status: str) -> None:
        await _db(client)[ATTEMPTS_COLLECTION].insert_one(
            {
                "competition_id": ObjectId(cid),
                "account_id": ObjectId(),
                "idempotency_key": f"{track}-{status}",
                "track": track,
                "status": status,
            }
        )

    asyncio.run(write_attempt("private", "FAILED"))
    asyncio.run(write_attempt("private", STATUS_QUEUED))
    login(client)

    private_stats = client.get(
        f"/api/admin/competitions/{cid}/submissions/stats?track=private"
    ).json()
    assert private_stats == {"total": 3, "completed": 1, "failed": 1, "in_flight": 1}
    public_stats = client.get(
        f"/api/admin/competitions/{cid}/submissions/stats?track=public"
    ).json()
    assert public_stats == {"total": 1, "completed": 1, "failed": 0, "in_flight": 0}


def test_scope_stats_bat_buoc_nhanh_voi_dual_va_tu_choi_nhanh_voi_single(client):
    login(client)
    dual = create_dual_ok(client, dual_body(slug="stats-scope-cup"))
    missing = client.get(f"/api/admin/competitions/{dual['id']}/submissions/stats")
    assert missing.status_code == 422
    assert err(missing)["code"] == "TRACK_REQUIRED"
    unknown = client.get(f"/api/admin/competitions/{dual['id']}/submissions/stats?track=beta")
    assert unknown.status_code == 422
    assert err(unknown)["code"] == "INVALID_TRACK"

    now = datetime.now(timezone.utc)
    single = client.post(
        "/api/admin/competitions",
        json={
            "slug": "stats-single-cup",
            "name": "stats-single-cup",
            "start_at": iso(now - timedelta(hours=1)),
            "end_at": iso(now + timedelta(hours=1)),
        },
    )
    assert single.status_code == 201, single.text
    rejected = client.get(
        f"/api/admin/competitions/{single.json()['id']}/submissions/stats?track=private"
    )
    assert rejected.status_code == 422
    assert err(rejected)["code"] == "INVALID_TRACK"


def test_immediate_policy_releases_on_publish(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="immediate-cup", result_policy="immediate"))
    cid = payload["id"]
    assert (
        client.get(f"/api/admin/competitions/{cid}").json()["tracks"]["private"]["results_released"]
        is False
    )
    published = publish_dual_v1(client, payload["id"], slug="immediate-cup")
    assert published["tracks"]["private"]["results_released"] is True
    assert published["tracks"]["private"]["results_published_at"] is not None


def test_switching_to_immediate_requires_reveal_confirmation(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="reveal-cup"))
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="reveal-cup")
    revision = client.get(f"/api/admin/competitions/{cid}").json()["control_revision"]

    unconfirmed = client.patch(
        f"/api/admin/competitions/{cid}/tracks/private/policy",
        json={"result_policy": "immediate", "expected_revision": revision, "reason": "hiện ngay"},
    )
    assert unconfirmed.status_code == 422

    confirmed = client.patch(
        f"/api/admin/competitions/{cid}/tracks/private/policy",
        json={
            "result_policy": "immediate",
            "expected_revision": revision,
            "reason": "BTC quyết định hiện ngay",
            "confirm_reveal": True,
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["tracks"]["private"]["result_policy"] == "immediate"
    assert confirmed.json()["tracks"]["private"]["results_released"] is True


def test_policy_patch_before_release_keeps_results_hidden(client):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="policy-cup"))
    cid = payload["id"]
    publish_dual_v1(client, payload["id"], slug="policy-cup")
    revision = client.get(f"/api/admin/competitions/{cid}").json()["control_revision"]
    changed = client.patch(
        f"/api/admin/competitions/{cid}/tracks/private/policy",
        json={
            "publish_condition": "after_closed_and_scored",
            "expected_revision": revision,
            "reason": "Siết điều kiện công bố",
        },
    )
    assert changed.status_code == 200, changed.text
    assert changed.json()["tracks"]["private"]["publish_condition"] == "after_closed_and_scored"
    assert changed.json()["tracks"]["private"]["results_released"] is False


# --- Clone cuộc thi dual ---------------------------------------------------------------------------


def test_clone_dual_copies_tracks_ground_truths_and_policy(client, isolated_data_dir):
    login(client)
    resource = {"label": "Đề bài", "url": "https://drive.google.com/file/d/clone-private"}
    payload = create_dual_ok(
        client,
        dual_body(
            slug="clone-cup",
            public_hours=(-5, -1),
            private_hours=(-3, 1),
            quota_per_day=9,
            result_policy="immediate",
            private_resources=[resource],
        ),
    )
    cid = payload["id"]
    publish_dual_v1(client, cid, slug="clone-cup")
    source = client.get(f"/api/admin/competitions/{cid}").json()

    cloned = client.post(f"/api/admin/competitions/{cid}/clone")
    assert cloned.status_code == 201, cloned.text
    clone = cloned.json()
    assert clone["mode"] == "public_private"
    assert clone["status"] == "draft"
    assert clone["control_revision"] == 1
    assert clone["stop_generation"] == 0
    assert clone["scoring_locked"] is False
    assert clone["publish_ready"] is True

    # Cửa sổ sớm nhất mở tại thời điểm clone; độ lệch và thời lượng của hai nhánh giữ nguyên.
    clone_public_start = at(clone["tracks"]["public"]["start_at"])
    clone_private_start = at(clone["tracks"]["private"]["start_at"])
    assert abs((clone_public_start - datetime.now(timezone.utc)).total_seconds()) < 30
    assert near(clone_private_start - clone_public_start, timedelta(hours=2))
    assert near(at(clone["tracks"]["public"]["end_at"]) - clone_public_start, timedelta(hours=4))
    assert near(at(clone["tracks"]["private"]["end_at"]) - clone_private_start, timedelta(hours=4))

    assert clone["tracks"]["public"]["quota_per_day"] == 9
    assert clone["tracks"]["private"]["quota_per_day"] == 9
    assert clone["tracks"]["private"]["resources"] == source["tracks"]["private"]["resources"]
    assert clone["tracks"]["private"]["result_policy"] == "immediate"
    # Chính sách được giữ nhưng dấu mốc công bố thì không: bản sao chỉ release khi publish lại.
    assert clone["tracks"]["private"]["results_released"] is False
    for track in ("public", "private"):
        assert (
            clone["tracks"][track]["ground_truth"]["sha256"]
            == source["tracks"][track]["ground_truth"]["sha256"]
        )
        assert clone["tracks"][track]["admission_seq"] == 0

    # GT của bản sao đọc được từ thư mục mới: publish chạy qua đúng cổng readiness.
    published = client.post(f"/api/admin/competitions/{clone['id']}/publish")
    assert published.status_code == 200, published.text
    assert published.json()["tracks"]["private"]["results_released"] is True


def test_clone_dual_v2_starts_without_copied_verification(client, fake_runner):
    login(client)
    payload = create_dual_ok(client, dual_body(slug="clone-v2-cup"))
    cid = payload["id"]
    assert (
        put_scoring_v2(client, cid, expected_revision=0, output_contract=V2_CONTRACT).status_code
        == 200
    )
    assert (
        upload_gt(client, cid, V2_GROUND_TRUTH, track="public", expected_revision=1).status_code
        == 200
    )
    assert (
        upload_gt(
            client, cid, flipped_v2_ground_truth(), track="private", expected_revision=2
        ).status_code
        == 200
    )
    assert run_test(client, cid, track="public", expected_revision=3).status_code == 200
    assert run_test(client, cid, track="private", expected_revision=3).status_code == 200
    assert client.post(f"/api/admin/competitions/{cid}/publish").status_code == 200
    source = client.get(f"/api/admin/competitions/{cid}").json()
    assert source["tracks"]["public"]["verified"] is True

    cloned = client.post(f"/api/admin/competitions/{cid}/clone")
    assert cloned.status_code == 201, cloned.text
    clone = cloned.json()
    # Cấu hình chấm và GT chép sang, nhưng thế hệ chạy thử thì không: bản sao phải tự xác minh lại.
    assert (
        clone["tracks"]["public"]["ground_truth"]["sha256"]
        == source["tracks"]["public"]["ground_truth"]["sha256"]
    )
    assert clone["tracks"]["public"]["verified"] is False
    assert clone["tracks"]["private"]["verified"] is False
    assert clone["publish_ready"] is False
    assert clone["publish_blocked_reason"]["code"] == "SCORING_TEST_REQUIRED"

    scoring = client.get(f"/api/admin/competitions/{clone['id']}/scoring").json()
    assert scoring["scoring"]["revision"] == 0
    assert scoring["scoring"]["evaluator"]["runtime_id"] is None
    assert scoring["tracks"]["public"]["verification"] is None

    for track in ("public", "private"):
        assert run_test(client, clone["id"], track=track, expected_revision=0).status_code == 200
    published = client.post(f"/api/admin/competitions/{clone['id']}/publish")
    assert published.status_code == 200, published.text


# --- Giai đoạn C: cổng admission, quota theo nhánh, vòng đời lượt ---------------------------------
#
# Thời gian được tua bằng cách sửa thẳng document (lượt, membership, cuộc thi), không bằng cách ngủ.
# Các lượt "chết giữa chừng" được dựng ở tầng service/store đúng như một tiến trình đứt tay, vì
# không có request HTTP nào tự tạo được trạng thái đó.

# Bài nộp hợp lệ cho cả hai nhánh v1: cùng bộ ID với hai ground truth của `publish_dual_v1`.
V1_SUBMISSION = b"id,prediction\n1,1\n2,1\n3,0\n4,0\n"


def _db(client):
    return client.app.state.mongo.db


def _patch(client, collection, document_id, patch: dict) -> None:
    async def run():
        await _db(client)[collection].update_one({"_id": ObjectId(document_id)}, patch)

    asyncio.run(run())


def _competition_doc(client, competition_id: str) -> dict:
    async def load():
        return await _db(client)[COMPETITIONS_COLLECTION].find_one(
            {"_id": ObjectId(competition_id)}
        )

    return asyncio.run(load())


def _attempt_doc(client, attempt_id) -> dict:
    async def load():
        return await _db(client)[ATTEMPTS_COLLECTION].find_one({"_id": ObjectId(attempt_id)})

    return asyncio.run(load())


def _out_of_time(client, attempt_id) -> None:
    """Tua lượt về quá khứ: hạn 60 giây của nó đã trôi qua."""
    _patch(
        client,
        ATTEMPTS_COLLECTION,
        attempt_id,
        {"$set": {"deadline_at": datetime.now(timezone.utc) - timedelta(seconds=1)}},
    )


def _reconcile(client) -> dict:
    return asyncio.run(
        worker.reconcile(_db(client), settings=get_settings(), now=datetime.now(timezone.utc))
    )


def _track_quota(client, competition_id: str, track: str) -> dict | None:
    """Bucket quota của một nhánh trên membership; None nghĩa là nhánh chưa từng tiêu suất nào."""
    membership = membership_document(client, competition_id)
    return (membership.get("track_quotas") or {}).get(track)


def _track_used(client, competition_id: str, track: str) -> int:
    return (_track_quota(client, competition_id, track) or {}).get("used", 0)


def _participant_context(client, competition_id: str) -> tuple[dict, dict, dict]:
    """Ba document nền của một lượt: cuộc thi, tài khoản thí sinh, membership."""

    async def load():
        db = _db(client)
        competition = await db[COMPETITIONS_COLLECTION].find_one(
            {"_id": ObjectId(competition_id)}
        )
        account = await db[ACCOUNTS_COLLECTION].find_one({"email": PARTICIPANT_CREDENTIALS[0]})
        membership = await db[MEMBERSHIPS_COLLECTION].find_one(
            {"competition_id": ObjectId(competition_id)}
        )
        return competition, account, membership

    return asyncio.run(load())


def publish_dual_v2(client, competition_id: str, *, slug: str) -> dict:
    """Đưa cuộc thi dual qua cổng publish bằng bộ chấm v2; hai nhánh có bằng chứng chạy thử riêng.

    Kết thúc ở phiên thí sinh để test nộp bài ngay được, giống `publish_v2_competition`.
    """
    assert (
        put_scoring_v2(
            client, competition_id, expected_revision=0, output_contract=V2_CONTRACT
        ).status_code
        == 200
    )
    assert (
        upload_gt(
            client, competition_id, V2_GROUND_TRUTH, track="public", expected_revision=1
        ).status_code
        == 200
    )
    assert (
        upload_gt(
            client,
            competition_id,
            flipped_v2_ground_truth(),
            track="private",
            expected_revision=2,
        ).status_code
        == 200
    )
    for track in ("public", "private"):
        assert (
            run_test(client, competition_id, track=track, expected_revision=3).status_code == 200
        )
    published = client.post(f"/api/admin/competitions/{competition_id}/publish")
    assert published.status_code == 200, published.text
    login_participant(client)
    assert client.post(f"/api/competitions/{slug}/join", json={}).status_code == 200
    return published.json()


def test_quota_moi_nhanh_giu_suat_rieng(client):
    """Mỗi nhánh tiêu suất của chính nó: Public hết lượt không đóng cửa Private và ngược lại."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="quota-track-cup", quota_per_day=1))
    publish_dual_v1(client, payload["id"], slug="quota-track-cup")
    cid = payload["id"]
    login_participant(client)

    first_public = submit(client, cid, V1_SUBMISSION, track="public")
    assert first_public.status_code == 201, first_public.text
    assert _track_used(client, cid, "public") == 1

    over_public = submit(client, cid, V1_SUBMISSION, track="public")
    assert over_public.status_code == 429
    assert err(over_public)["code"] == "SUBMISSION_QUOTA_EXCEEDED"

    first_private = submit(client, cid, V1_SUBMISSION, track="private")
    assert first_private.status_code == 201, first_private.text
    assert _track_used(client, cid, "private") == 1

    over_private = submit(client, cid, V1_SUBMISSION, track="private")
    assert over_private.status_code == 429
    assert err(over_private)["code"] == "SUBMISSION_QUOTA_EXCEEDED"

    submissions = submission_documents(client)
    assert sorted(item["track"] for item in submissions) == ["private", "public"]
    # Mốc nhận bài của từng nhánh được ghim vào chính bài nộp.
    for item in submissions:
        assert item["admitted_at"] is not None
        assert item["admitted_end_at"] is not None


def test_suat_quota_theo_nhanh_nguyen_tu_khi_nop_dong_thoi(client):
    """Cổng chặn thật là `$inc` có điều kiện trên bucket của đúng nhánh, không phải phép đếm."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="quota-race-cup", quota_per_day=5))
    publish_dual_v1(client, payload["id"], slug="quota-race-cup")
    cid = payload["id"]
    membership = membership_document(client, cid)

    async def reserve():
        return await submissions_service.reserve_quota_slot(
            _db(client), membership, 5, datetime.now(timezone.utc), track="public"
        )

    async def race():
        return await asyncio.gather(*[reserve() for _ in range(8)])

    results = asyncio.run(race())
    assert len([item for item in results if item is not None]) == 5
    assert _track_used(client, cid, "public") == 5
    # Nhánh Private chưa ai đụng tới: bucket của nó không được tạo.
    assert _track_quota(client, cid, "private") is None


def test_cong_admission_dong_cua_thi_hoan_suat(client):
    """Cửa sổ đã qua tại đúng thao tác ghi: lượt bị đóng ngay, suất vừa giữ được hoàn, không cấp số."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="gate-closed-cup"))
    publish_dual_v1(client, payload["id"], slug="gate-closed-cup")
    cid = payload["id"]

    # Thí sinh bấm Nút lúc cửa còn mở, nhưng tới lúc ghi thì đồng hồ đã qua `end_at` của nhánh.
    _patch(
        client,
        COMPETITIONS_COLLECTION,
        cid,
        {"$set": {"tracks.private.end_at": datetime.now(timezone.utc) - timedelta(seconds=1)}},
    )
    competition, account, membership = _participant_context(client, cid)
    now = datetime.now(timezone.utc)

    with pytest.raises(HTTPException) as caught:
        asyncio.run(
            attempts_service.admit_inline(
                _db(client),
                competition=competition,
                account=account,
                membership=membership,
                data=V1_SUBMISSION,
                notebook_data=VALID_NOTEBOOK,
                csv_filename="answers.csv",
                notebook_filename="solution.ipynb",
                idempotency_key="qua-cong-admission",
                received_at=now,
                now=now,
                track="private",
            )
        )
    assert caught.value.status_code == 422
    assert caught.value.detail["code"] == "SUBMISSION_DEADLINE_PASSED"

    attempts = attempt_documents(client)
    assert [item["status"] for item in attempts] == ["FAILED"]
    assert attempts[0]["error"]["code"] == "SUBMISSION_DEADLINE_PASSED"
    # Suất đã hoàn và cổng không cấp số nào: lượt chết trước khi thành bài.
    assert _track_used(client, cid, "private") == 0
    assert int(_competition_doc(client, cid)["tracks"]["private"].get("admission_seq", 0)) == 0
    assert submission_documents(client) == []


def test_nop_sau_han_mot_nhanh_khong_ton_suat_nhanh_kia(client):
    """Upload xong sau `end_at` bị cổng từ chối; nhánh còn mở vẫn nhận bài bình thường."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="late-track-cup"))
    publish_dual_v1(client, payload["id"], slug="late-track-cup")
    cid = payload["id"]
    _patch(
        client,
        COMPETITIONS_COLLECTION,
        cid,
        {"$set": {"tracks.private.end_at": datetime.now(timezone.utc) - timedelta(seconds=1)}},
    )
    login_participant(client)

    late = submit(client, cid, V1_SUBMISSION, track="private")
    assert late.status_code == 422
    assert err(late)["code"] == "SUBMISSION_DEADLINE_PASSED"
    assert attempt_documents(client) == []
    assert _track_quota(client, cid, "private") is None

    # Cửa sổ là chuyện của từng nhánh: Public vẫn nộp được.
    assert submit(client, cid, V1_SUBMISSION, track="public").status_code == 201


def test_luot_admission_truoc_han_van_duoc_cham_sau_han(client, fake_runner):
    """Lượt đã admission hợp lệ sống qua hạn nhánh: cửa sổ ghim trong intent mới là thứ quyết định."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="after-end-cup"))
    publish_dual_v2(client, payload["id"], slug="after-end-cup")
    cid = payload["id"]
    original_end = _competition_doc(client, cid)["tracks"]["public"]["end_at"]

    queued = submit(client, cid, V2_SUBMISSION, track="public")
    assert queued.status_code == 202, queued.text

    # Đồng hồ trôi qua hạn Public trong lúc lượt còn chờ: worker không được từ chối vì `now > end_at`.
    _patch(
        client,
        COMPETITIONS_COLLECTION,
        cid,
        {"$set": {"tracks.public.end_at": datetime.now(timezone.utc) - timedelta(seconds=1)}},
    )
    assert run_worker(client) == 1

    body = attempt_status(client, cid, queued.json()["attempt_id"])
    assert body["status"] == "COMPLETED"
    submission = submission_documents(client)[0]
    assert submission["track"] == "public"
    # Mốc nhận bài là mốc của chính lượt, không phải mốc sau khi admin dời lịch.
    assert submission["admitted_end_at"] == original_end
    assert submission["admitted_at"] <= submission["admitted_end_at"]
    assert _track_used(client, cid, "public") == 1


def test_luot_cho_chet_truoc_khi_vao_hang_doi_thi_tra_cho_va_hoan_suat(client, fake_runner):
    """Đứt tay giữa lúc giữ quota và lúc file vào kho tạm: đối soát đóng lượt, trả chỗ, hoàn một suất."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="crash-queue-cup"))
    publish_dual_v2(client, payload["id"], slug="crash-queue-cup")
    cid = payload["id"]
    competition, account, membership = _participant_context(client, cid)
    now = datetime.now(timezone.utc)
    attempt_id = ObjectId()

    async def crash():
        db = _db(client)
        used = await submissions_service.reserve_quota_slot(
            db, membership, 5, now, track="public", attempt_id=str(attempt_id)
        )
        assert used == 1
        return await attempts_store.admit(
            db,
            attempt_id=attempt_id,
            competition_id=competition["_id"],
            account_id=account["_id"],
            membership_id=membership["_id"],
            idempotency_key="crash-queue",
            payload_sha256="0" * 64,
            staging_prefix="staging/scoring/crash/",
            deadline_at=now + timedelta(seconds=60),
            now=now,
            capacity=20,
            track="public",
        )

    attempt = asyncio.run(crash())
    assert attempt["status"] == STATUS_STAGING
    assert "queue_slot" in attempt

    _out_of_time(client, attempt_id)
    assert _reconcile(client)["expired"] == 1
    assert _reconcile(client) == {"expired": 0, "requeued": 0, "resolved": 0, "refunded": 0}

    closed = _attempt_doc(client, attempt_id)
    assert closed["status"] == "EXPIRED"
    assert closed["error"]["code"] == "SUBMISSION_EXPIRED"
    assert "queue_slot" not in closed
    assert _track_used(client, cid, "public") == 0
    assert submission_documents(client) == []
    assert run_worker(client) == 0


def test_vo_inline_chet_truoc_khi_ghi_thi_dong_luot_va_hoan_suat(client):
    """Request inline đứt sau khi qua cổng admission: lượt không thành bài và không bị chấm lại."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="crash-inline-cup"))
    publish_dual_v1(client, payload["id"], slug="crash-inline-cup")
    cid = payload["id"]
    competition, account, membership = _participant_context(client, cid)
    now = datetime.now(timezone.utc)

    attempt, replayed = asyncio.run(
        attempts_service.admit_inline(
            _db(client),
            competition=competition,
            account=account,
            membership=membership,
            data=V1_SUBMISSION,
            notebook_data=VALID_NOTEBOOK,
            csv_filename="answers.csv",
            notebook_filename="solution.ipynb",
            idempotency_key="crash-inline",
            received_at=now,
            now=now,
            track="private",
        )
    )
    assert replayed is False
    assert attempt["status"] == STATUS_STAGING
    assert attempt["admission"]["admission_seq"] == 1
    assert _track_used(client, cid, "private") == 1

    # Tiến trình chết ngay sau đó: không begin, không ghi bài. Hết 60 giây, đối soát dọn.
    _out_of_time(client, attempt["_id"])
    assert _reconcile(client)["expired"] == 1
    assert _reconcile(client) == {"expired": 0, "requeued": 0, "resolved": 0, "refunded": 0}

    closed = _attempt_doc(client, attempt["_id"])
    assert closed["status"] == "EXPIRED"
    assert closed["error"]["code"] == "SUBMISSION_EXPIRED"
    assert _track_used(client, cid, "private") == 0
    # Lượt inline không có ai chấm lại: worker không thấy gì để nhận.
    assert run_worker(client) == 0
    assert submission_documents(client) == []


def test_vo_inline_bo_dang_thi_dong_luot_va_khong_hoi_sinh(client):
    """Request chết giữa lúc chấm inline: lease hết hạn là lượt hỏng, không quay lại hàng đợi."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="abandoned-inline-cup"))
    publish_dual_v1(client, payload["id"], slug="abandoned-inline-cup")
    cid = payload["id"]
    competition, account, membership = _participant_context(client, cid)
    now = datetime.now(timezone.utc)

    async def admit_and_begin():
        db = _db(client)
        attempt, _replayed = await attempts_service.admit_inline(
            db,
            competition=competition,
            account=account,
            membership=membership,
            data=V1_SUBMISSION,
            notebook_data=VALID_NOTEBOOK,
            csv_filename="answers.csv",
            notebook_filename="solution.ipynb",
            idempotency_key="bo-dang-inline",
            received_at=now,
            now=now,
            track="public",
        )
        return attempt, await attempts_store.begin_inline(
            db, attempt, now=now, lease_seconds=120
        )

    attempt, started = asyncio.run(admit_and_begin())
    assert started["status"] == "RUNNING"

    _patch(
        client,
        ATTEMPTS_COLLECTION,
        attempt["_id"],
        {"$set": {"lease_expires_at": datetime.now(timezone.utc) - timedelta(seconds=1)}},
    )
    assert _reconcile(client)["expired"] == 1

    closed = _attempt_doc(client, attempt["_id"])
    assert closed["status"] == "EXPIRED"
    assert closed["error"]["code"] == "SUBMISSION_EXPIRED"
    assert closed["lease_token"] is None
    assert _track_used(client, cid, "public") == 0
    # Lượt inline không bao giờ được requeue, nên không có lần chấm thứ hai nào.
    assert _reconcile(client) == {"expired": 0, "requeued": 0, "resolved": 0, "refunded": 0}
    assert run_worker(client) == 0


def test_lenh_dong_huy_luot_cu_va_mo_lai_khong_hoi_sinh(client, fake_runner):
    """Generation của lệnh đóng: lượt nhận theo thế hệ cũ bị hủy, mở lại không làm nó sống dậy."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="generation-cup"))
    publish_dual_v2(client, payload["id"], slug="generation-cup")
    cid = payload["id"]

    queued = submit(client, cid, V2_SUBMISSION, track="private")
    assert queued.status_code == 202, queued.text

    login(client)
    assert client.post(f"/api/admin/competitions/{cid}/close").status_code == 200
    assert client.post(f"/api/admin/competitions/{cid}/reopen", json={}).status_code == 200

    assert run_worker(client) == 1
    login_participant(client)
    body = attempt_status(client, cid, queued.json()["attempt_id"])
    assert body["status"] == "FAILED"
    assert body["error"] == {
        "code": attempts_service.CANCELLED_CODE,
        "message": attempts_service.CANCELLED_MESSAGE,
    }
    assert submission_documents(client) == []
    assert _track_used(client, cid, "private") == 0

    # Mở lại không hồi sinh lượt đã hủy...
    assert run_worker(client) == 0
    # ...nhưng lượt mới theo thế hệ hiện tại vẫn được chấm bình thường.
    again = submit(client, cid, V2_SUBMISSION, track="private")
    assert again.status_code == 202, again.text
    assert run_worker(client) == 1
    assert attempt_status(client, cid, again.json()["attempt_id"])["status"] == "COMPLETED"
    assert _track_used(client, cid, "private") == 1


def test_idempotency_dual_tra_lai_dung_luot_va_track_khac_la_xung_dot(client, fake_runner):
    """Một lần nhấn Nút chỉ có một lượt: gửi lại nhận đúng lượt cũ; đổi nhánh hay đổi file là 409."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="idem-dual-cup"))
    publish_dual_v2(client, payload["id"], slug="idem-dual-cup")
    cid = payload["id"]

    first = submit(client, cid, V2_SUBMISSION, track="public", key="dual-key-1")
    again = submit(client, cid, V2_SUBMISSION, track="public", key="dual-key-1")
    assert (first.status_code, again.status_code) == (202, 202)
    assert again.json()["attempt_id"] == first.json()["attempt_id"]
    assert len(attempt_documents(client)) == 1
    assert _track_used(client, cid, "public") == 1

    # Cùng key mà đổi nhánh: một lần nhấn không thể vừa Public vừa Private.
    other_track = submit(client, cid, V2_SUBMISSION, track="private", key="dual-key-1")
    assert other_track.status_code == 409
    assert err(other_track)["code"] == "IDEMPOTENCY_CONFLICT"

    other_file = submit(
        client, cid, b"id,predict_type\n1,B\n2,A\n3,B\n4,A\n", track="public", key="dual-key-1"
    )
    assert other_file.status_code == 409
    assert err(other_file)["code"] == "IDEMPOTENCY_CONFLICT"
    assert len(attempt_documents(client)) == 1
    assert _track_quota(client, cid, "private") is None


def test_hoan_suat_khong_ro_ri_sang_ngay_khac(client, fake_runner):
    """Lượt của ngày cũ không trừ vào bộ đếm ngày mới; ngày mới tự seed lại từ bài đã hoàn tất."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="rollover-cup"))
    publish_dual_v2(client, payload["id"], slug="rollover-cup")
    cid = payload["id"]

    queued = submit(client, cid, V2_SUBMISSION, track="public")
    assert queued.status_code == 202, queued.text
    attempt_id = queued.json()["attempt_id"]
    yesterday = utc_day_key(datetime.now(timezone.utc) - timedelta(days=1))

    # Đồng hồ sang ngày mới: bộ đếm còn nguyên dấu của lượt hôm qua.
    _patch(
        client,
        MEMBERSHIPS_COLLECTION,
        membership_document(client, cid)["_id"],
        {
            "$set": {
                "track_quotas.public.day": yesterday,
                "track_quotas.public.used": 1,
                "track_quotas.public.claims": {attempt_id: yesterday},
            }
        },
    )

    _out_of_time(client, attempt_id)
    assert _reconcile(client)["expired"] == 1
    # Hoàn suất là chuyện của đúng ngày đã giữ: bộ đếm hôm qua không bị trừ.
    bucket = _track_quota(client, cid, "public")
    assert (bucket["day"], bucket["used"]) == (yesterday, 1)

    # Ngày mới seed lại từ số bài đã hoàn tất (hôm nay chưa có bài nào) nên lượt mới vào được.
    again = submit(client, cid, V2_SUBMISSION, track="public")
    assert again.status_code == 202, again.text
    bucket = _track_quota(client, cid, "public")
    assert bucket["day"] == utc_day_key(datetime.now(timezone.utc))
    assert bucket["used"] == 1


def test_idempotency_dual_inline_tra_lai_dung_bai_da_ghi(client):
    """Đường inline: gửi lại cùng key nhận đúng bài đã chấm, không chấm lần hai."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="idem-inline-cup"))
    publish_dual_v1(client, payload["id"], slug="idem-inline-cup")
    cid = payload["id"]
    login_participant(client)

    first = submit(client, cid, V1_SUBMISSION, track="public", key="inline-key-1")
    assert first.status_code == 201, first.text
    again = submit(client, cid, V1_SUBMISSION, track="public", key="inline-key-1")
    assert again.status_code == 201, again.text
    assert again.json()["id"] == first.json()["id"]
    assert len(submission_documents(client)) == 1
    assert _track_used(client, cid, "public") == 1


def test_danh_sach_luot_cho_loc_duoc_theo_nhanh_va_bien_nhan_lo_dung_nhanh(client, fake_runner):
    """Trang nộp bài dual mở lại theo một nhánh: lượt chờ phải lọc và khai đúng nhánh."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="attempt-track-cup"))
    cid = payload["id"]
    publish_dual_v2(client, cid, slug="attempt-track-cup")
    posted = submit(client, cid, V2_SUBMISSION, track="private")
    assert posted.status_code == 202, posted.text
    attempt_id = posted.json()["attempt_id"]
    # Biên nhận nói đúng nhánh ngay từ response nhận bài và ở mọi lần hỏi trạng thái sau đó.
    assert posted.json()["track"] == "private"
    assert attempt_status(client, cid, attempt_id)["track"] == "private"

    assert [
        row["attempt_id"]
        for row in client.get(f"/api/competitions/{cid}/submissions/attempts").json()["attempts"]
    ] == [attempt_id]
    private_only = client.get(
        f"/api/competitions/{cid}/submissions/attempts", params={"track": "private"}
    )
    assert [row["attempt_id"] for row in private_only.json()["attempts"]] == [attempt_id]
    public_only = client.get(
        f"/api/competitions/{cid}/submissions/attempts", params={"track": "public"}
    )
    assert public_only.json()["attempts"] == []

    bad = client.get(f"/api/competitions/{cid}/submissions/attempts", params={"track": "beta"})
    assert bad.status_code == 422
    assert err(bad)["code"] == "INVALID_TRACK"


def test_danh_sach_luot_cho_cua_single_khong_co_khoa_track_va_tu_choi_loc(client, fake_runner):
    """Single giữ nguyên payload cũ: không khóa track, và lọc theo track là lỗi hợp đồng."""
    competition = publish_v2_competition(client, slug="attempt-single-cup")
    cid = competition["id"]
    posted = submit(client, cid, V2_SUBMISSION)
    assert posted.status_code == 202, posted.text
    assert "track" not in posted.json()

    listed = client.get(f"/api/competitions/{cid}/submissions/attempts").json()["attempts"]
    assert [row["attempt_id"] for row in listed] == [posted.json()["attempt_id"]]
    assert all("track" not in row for row in listed)

    rejected = client.get(
        f"/api/competitions/{cid}/submissions/attempts", params={"track": "public"}
    )
    assert rejected.status_code == 422
    assert err(rejected)["code"] == "INVALID_TRACK"


# --- Giai đoạn D: công bố và riêng tư ---------------------------------------------------------------
#
# Gate của giai đoạn D: giá trị Private chưa công bố không được rời backend ở bất kỳ payload nào của
# thí sinh - không chỉ bị che trên giao diện. Vì vậy ngoài việc kiểm cấu trúc DTO, các test dưới đây
# quét thẳng chữ trong response để tìm giá trị điểm nhận dạng được.

# Hai giá trị nhận dạng được: có mặt trong payload nghĩa là điểm đã rời backend.
PRIVATE_SCORE = 0.123456789
PUBLIC_SCORE = 0.987654321


def _release_private(client, competition_id: str) -> dict:
    """Công bố kết quả Private bằng quyền admin rồi trả về phiên thí sinh."""
    login(client)
    revision = client.get(f"/api/admin/competitions/{competition_id}").json()["control_revision"]
    released = client.post(
        f"/api/admin/competitions/{competition_id}/tracks/private/publish-results",
        json={"expected_revision": revision, "reason": "Công bố kết quả Private"},
    )
    assert released.status_code == 200, released.text
    login_participant(client)
    return released.json()


def _private_submission(client, fake_runner, slug: str, score: float) -> tuple[str, str, dict]:
    """Cuộc thi dual v2 đã publish với một bài Private đã chấm xong nhưng chưa công bố.

    Trả `(competition_id, attempt_id, response 202 của lần nộp)`; kết thúc ở phiên thí sinh.
    """
    login(client)
    payload = create_dual_ok(client, dual_body(slug=slug))
    publish_dual_v2(client, payload["id"], slug=slug)
    fake_runner.metrics = {"accuracy": score, "n_items": 4.0}
    posted = submit(client, payload["id"], V2_SUBMISSION, track="private")
    assert posted.status_code == 202, posted.text
    attempt_id = posted.json()["attempt_id"]
    assert run_worker(client) == 1
    return payload["id"], attempt_id, posted


def test_diem_private_chua_cong_bo_khong_ro_ri_o_bat_ky_bay_nao(client, fake_runner):
    """Bài Private chưa công bố: mọi bề mặt thí sinh giữ trạng thái chấm nhưng không mang điểm."""
    cid, attempt_id, posted = _private_submission(client, fake_runner, "privacy-cup", PRIVATE_SCORE)
    secret = repr(PRIVATE_SCORE)

    poll = client.get(f"/api/competitions/{cid}/submissions/attempts/{attempt_id}")
    assert poll.status_code == 200, poll.text
    polled = poll.json()["submission"]
    assert polled["status"] == "completed"
    assert polled["result_visibility"] == "hidden"
    assert polled["visibility_reason"] == "private_unpublished"
    assert polled["primary_score"] is None
    assert polled["metrics"] == {}
    assert polled["track"] == "private"
    assert "normalization_snapshot" not in polled

    history = client.get(f"/api/competitions/{cid}/submissions/me")
    item = history.json()["submissions"][0]
    assert item["result_visibility"] == "hidden"
    assert item["primary_score"] is None
    assert item["metrics"] == {}
    assert item["track"] == "private"

    board = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "private"})
    assert board.status_code == 403
    assert err(board)["code"] == "PRIVATE_RESULTS_UNPUBLISHED"

    listing = client.get("/api/competitions")
    card = next(c for c in listing.json()["competitions"] if c["id"] == cid)
    assert set(card["my_stats_by_track"]) == {"public", "private"}
    private_stats = card["my_stats_by_track"]["private"]
    assert private_stats["rank"] is None
    assert private_stats["best_score"] is None
    assert private_stats["best_normalized_score"] is None
    # Quota là số thật của chính thí sinh: chờ công bố không làm nó biến mất.
    assert private_stats["used_today"] == 1
    assert "my_stats" not in card

    detail = client.get("/api/competitions/privacy-cup")
    tracks = detail.json()["tracks"]
    assert tracks["private"]["results_released"] is False
    # Quota từng nhánh là số thật: chờ công bố không làm nó biến mất, và Public chưa tiêu suất nào.
    assert tracks["private"]["quota"]["used_today"] == 1
    assert tracks["private"]["quota"]["remaining"] == 4
    assert tracks["public"]["quota"]["used_today"] == 0

    for response in (posted, poll, history, board, listing, detail):
        assert secret not in response.text

    # Cùng tài khoản, nhánh Public vẫn hiện bình thường: gate theo nhánh, không theo người.
    fake_runner.metrics = {"accuracy": PUBLIC_SCORE, "n_items": 4.0}
    public_posted = submit(client, cid, V2_SUBMISSION, track="public")
    assert public_posted.status_code == 202, public_posted.text
    assert run_worker(client) == 1
    public_submission = attempt_status(client, cid, public_posted.json()["attempt_id"])["submission"]
    assert public_submission["result_visibility"] == "visible"
    assert public_submission["primary_score"] == PUBLIC_SCORE

    both = client.get(f"/api/competitions/{cid}/submissions/me")
    by_track = {row["track"]: row for row in both.json()["submissions"]}
    assert by_track["public"]["primary_score"] == PUBLIC_SCORE
    assert by_track["private"]["primary_score"] is None
    assert repr(PUBLIC_SCORE) in both.text
    assert secret not in both.text


def test_cong_bo_khong_cham_lai_va_mo_dung_bai_da_cham(client, fake_runner):
    """Công bố chỉ mở cửa: cùng bản ghi, cùng điểm, không có lần chấm thứ hai."""
    cid, attempt_id, _ = _private_submission(client, fake_runner, "reveal-cup", PRIVATE_SCORE)
    stored_before = submission_documents(client)[0]
    assert attempt_status(client, cid, attempt_id)["submission"]["result_visibility"] == "hidden"

    released = _release_private(client, cid)
    assert released["tracks"]["private"]["results_released"] is True

    revealed = attempt_status(client, cid, attempt_id)["submission"]
    assert revealed["result_visibility"] == "visible"
    assert revealed["primary_score"] == PRIVATE_SCORE
    assert revealed["metrics"]["accuracy"] == PRIVATE_SCORE
    assert "visibility_reason" not in revealed

    stored_after = submission_documents(client)[0]
    assert stored_after["primary_score"] == stored_before["primary_score"] == PRIVATE_SCORE
    assert stored_after["metrics"] == stored_before["metrics"]

    board = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "private"})
    assert board.status_code == 200, board.text
    assert board.json()["track"] == "private"
    assert board.json()["total"] == 1
    assert board.json()["entries"][0]["primary_score"] == PRIVATE_SCORE
    assert board.json()["me"]["is_current_user"] is True
    assert board.json()["me"]["rank"] == 1

    card = next(
        c for c in client.get("/api/competitions").json()["competitions"] if c["id"] == cid
    )
    assert card["my_stats_by_track"]["private"]["best_score"] == PRIVATE_SCORE


def test_bxh_dual_bat_buoc_chon_nhanh_va_khong_dung_chung_cache(client, fake_runner):
    """Hai bảng là hai population: thiếu nhánh bị từ chối, bảng đã cache của nhánh kia không lấn."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="board-cup"))
    cid = payload["id"]
    publish_dual_v2(client, cid, slug="board-cup")

    # Chấm từng nhánh ngay sau khi nộp: runner giả trả metric hiện tại của nó, nên hai bài phải
    # đi qua worker ở hai thời điểm khác nhau mới mang hai điểm khác nhau.
    fake_runner.metrics = {"accuracy": PUBLIC_SCORE, "n_items": 4.0}
    assert submit(client, cid, V2_SUBMISSION, track="public").status_code == 202
    assert run_worker(client) == 1
    fake_runner.metrics = {"accuracy": PRIVATE_SCORE, "n_items": 4.0}
    assert submit(client, cid, V2_SUBMISSION, track="private").status_code == 202
    assert run_worker(client) == 1

    missing = client.get(f"/api/competitions/{cid}/leaderboard")
    assert missing.status_code == 422
    assert err(missing)["code"] == "TRACK_REQUIRED"
    bad = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "side"})
    assert bad.status_code == 422
    assert err(bad)["code"] == "INVALID_TRACK"

    public = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "public"})
    assert public.status_code == 200, public.text
    assert public.json()["total"] == 1
    assert public.json()["entries"][0]["primary_score"] == PUBLIC_SCORE

    hidden = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "private"})
    assert hidden.status_code == 403
    assert err(hidden)["code"] == "PRIVATE_RESULTS_UNPUBLISHED"

    # Bảng Public vừa được cache; bảng Private sau công bố phải đọc đúng population của nó.
    _release_private(client, cid)
    private = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "private"})
    assert private.status_code == 200, private.text
    assert private.json()["total"] == 1
    assert private.json()["entries"][0]["primary_score"] == PRIVATE_SCORE


def test_master_bxh_an_khong_che_diem_da_cong_bo_va_nguoc_lai(client, fake_runner):
    """Master ẩn BXH và chưa công bố Private là hai luật độc lập, không suy ra nhau."""
    cid, attempt_id, _ = _private_submission(client, fake_runner, "master-cup", PRIVATE_SCORE)
    login(client)
    off = client.patch(
        f"/api/admin/competitions/{cid}", json={"leaderboard_visible": False}
    )
    assert off.status_code == 200, off.text
    login_participant(client)

    # Chưa công bố + master tắt: mã chặn vẫn là chuyện công bố, không phải chuyện ẩn bảng.
    hidden = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "private"})
    assert hidden.status_code == 403
    assert err(hidden)["code"] == "PRIVATE_RESULTS_UNPUBLISHED"

    _release_private(client, cid)
    # Đã công bố nhưng master tắt: BXH bị ẩn, còn điểm của chính thí sinh thì vẫn hiện.
    masked = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "private"})
    assert masked.status_code == 403
    assert err(masked)["code"] == "LEADERBOARD_HIDDEN"
    revealed = attempt_status(client, cid, attempt_id)["submission"]
    assert revealed["result_visibility"] == "visible"
    assert revealed["primary_score"] == PRIVATE_SCORE

    card = next(
        c for c in client.get("/api/competitions").json()["competitions"] if c["id"] == cid
    )
    assert card["my_stats_by_track"]["private"]["best_score"] is None
    assert card["my_stats_by_track"]["private"]["used_today"] == 1


def test_xuat_ket_qua_dual_bat_buoc_chon_nhanh_va_ghi_ro_trang_thai(client, fake_runner):
    """Export dual phải có scope; bản Private chưa công bố tự nhận là bản tạm ở tên file và sheet."""
    cid, _, _ = _private_submission(client, fake_runner, "export-cup", PRIVATE_SCORE)
    login(client)

    missing = client.get(f"/api/admin/competitions/{cid}/export.xlsx")
    assert missing.status_code == 422
    assert err(missing)["code"] == "TRACK_REQUIRED"
    bad = client.get(f"/api/admin/competitions/{cid}/export.xlsx", params={"track": "side"})
    assert bad.status_code == 422
    assert err(bad)["code"] == "INVALID_TRACK"

    provisional = client.get(
        f"/api/admin/competitions/{cid}/export.xlsx", params={"track": "private"}
    )
    assert provisional.status_code == 200, provisional.text
    assert "-private-provisional-results-" in provisional.headers["content-disposition"]
    sheet = load_workbook(BytesIO(provisional.content)).active
    assert sheet.title == "Results private (provisional)"
    # Admin xem đủ: điểm Private nằm trong file dù chưa công bố.
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert PRIVATE_SCORE in values

    _release_private(client, cid)
    login(client)
    released = client.get(
        f"/api/admin/competitions/{cid}/export.xlsx", params={"track": "private"}
    )
    assert released.status_code == 200, released.text
    assert "-private-results-" in released.headers["content-disposition"]
    assert "provisional" not in released.headers["content-disposition"]
    assert load_workbook(BytesIO(released.content)).active.title == "Results private"


def test_bang_admin_loc_duoc_theo_nhanh_va_luon_thay_du_diem(client, fake_runner):
    """Bảng admin lọc nhánh khi được hỏi; không lọc thì thấy cả hai nhánh kèm nhãn track."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="admin-list-cup"))
    cid = payload["id"]
    publish_dual_v2(client, cid, slug="admin-list-cup")
    fake_runner.metrics = {"accuracy": PUBLIC_SCORE, "n_items": 4.0}
    assert submit(client, cid, V2_SUBMISSION, track="public").status_code == 202
    assert run_worker(client) == 1
    fake_runner.metrics = {"accuracy": PRIVATE_SCORE, "n_items": 4.0}
    assert submit(client, cid, V2_SUBMISSION, track="private").status_code == 202
    assert run_worker(client) == 1

    login(client)
    both = client.get(f"/api/admin/competitions/{cid}/submissions")
    assert both.status_code == 200, both.text
    rows = {row["track"]: row for row in both.json()["submissions"]}
    assert set(rows) == {"public", "private"}
    # Admin xem đủ qua quyền admin thật: bài Private chưa công bố vẫn có điểm và không mang cờ che.
    assert rows["private"]["primary_score"] == PRIVATE_SCORE
    assert "result_visibility" not in rows["private"]

    only_private = client.get(
        f"/api/admin/competitions/{cid}/submissions", params={"track": "private"}
    )
    assert only_private.status_code == 200
    assert only_private.json()["total"] == 1
    assert only_private.json()["submissions"][0]["track"] == "private"

    bad = client.get(
        f"/api/admin/competitions/{cid}/submissions", params={"track": "side"}
    )
    assert bad.status_code == 422
    assert err(bad)["code"] == "INVALID_TRACK"


def test_luot_dang_cham_khi_cong_bo_thi_hoan_tat_sau_do_duoc_hien(client, fake_runner):
    """Publish in-flight: lượt nhận trước công bố nhưng ghi bài sau đó vẫn phải hiện."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="inflight-cup"))
    cid = payload["id"]
    publish_dual_v2(client, cid, slug="inflight-cup")
    fake_runner.metrics = {"accuracy": PRIVATE_SCORE, "n_items": 4.0}
    posted = submit(client, cid, V2_SUBMISSION, track="private")
    assert posted.status_code == 202, posted.text
    assert posted.json()["status"] == STATUS_QUEUED
    attempt_id = posted.json()["attempt_id"]

    _release_private(client, cid)
    assert run_worker(client) == 1
    submission = attempt_status(client, cid, attempt_id)["submission"]
    assert submission["result_visibility"] == "visible"
    assert submission["primary_score"] == PRIVATE_SCORE


def test_snapshot_chuan_hoa_cua_private_chua_cong_bo_khong_lo_diem(client):
    """Snapshot chuẩn hóa cũng là dữ liệu suy ra điểm: chưa công bố thì không rời backend."""
    login(client)
    body = dual_body(slug="norm-privacy")
    body["normalization"] = {"enabled": True, "baseline": 0.5}
    payload = create_dual_ok(client, body)
    cid = payload["id"]
    publish_dual_v1(client, cid, slug="norm-privacy")
    login_participant(client)

    submitted = submit(client, cid, V1_SUBMISSION, track="private")
    assert submitted.status_code == 201, submitted.text
    hidden = submitted.json()
    assert hidden["result_visibility"] == "hidden"
    assert hidden["primary_score"] is None
    assert hidden["metrics"] == {}
    assert "normalization_snapshot" not in hidden
    # Bản ghi thật vẫn có snapshot: công bố chỉ mở cửa, không dựng lại dữ liệu.
    assert submission_documents(client, {"track": "private"})[0]["normalization_snapshot"]

    _release_private(client, cid)
    item = client.get(f"/api/competitions/{cid}/submissions/me").json()["submissions"][0]
    assert item["result_visibility"] == "visible"
    assert item["normalization_snapshot"] is not None
    assert item["primary_score"] is not None


def test_tu_choi_bai_private_sau_cong_bo_chi_doi_bang_private(client, fake_runner):
    """Reject/restore sau công bố chỉ đổi đúng nhánh: bảng Public không được đụng tới."""
    login(client)
    payload = create_dual_ok(client, dual_body(slug="review-cup"))
    cid = payload["id"]
    publish_dual_v2(client, cid, slug="review-cup")
    fake_runner.metrics = {"accuracy": PUBLIC_SCORE, "n_items": 4.0}
    assert submit(client, cid, V2_SUBMISSION, track="public").status_code == 202
    assert run_worker(client) == 1
    fake_runner.metrics = {"accuracy": PRIVATE_SCORE, "n_items": 4.0}
    assert submit(client, cid, V2_SUBMISSION, track="private").status_code == 202
    assert run_worker(client) == 1
    private_submission_id = next(
        str(row["_id"]) for row in submission_documents(client, {"track": "private"})
    )
    _release_private(client, cid)

    login(client)
    rejected = client.patch(
        f"/api/admin/submissions/{private_submission_id}/review",
        json={"status": "rejected", "note": "Nghi vấn sao chép"},
    )
    assert rejected.status_code == 200, rejected.text

    private = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "private"})
    assert private.status_code == 200
    assert private.json()["total"] == 0
    login_participant(client)
    public = client.get(f"/api/competitions/{cid}/leaderboard", params={"track": "public"})
    assert public.status_code == 200
    assert public.json()["total"] == 1
    assert public.json()["entries"][0]["primary_score"] == PUBLIC_SCORE


def test_single_khong_doi_hinh_dang_dto_sau_khi_them_nhanh(client):
    """Single không có cờ công bố, không có my_stats_by_track và export giữ tên cũ."""
    competition = ready_competition(client, slug="single-shape")
    cid = competition["id"]

    submitted = submit(client, cid, V1_SUBMISSION)
    assert submitted.status_code == 201, submitted.text
    body = submitted.json()
    assert body["metrics"]
    assert "result_visibility" not in body
    assert "visibility_reason" not in body
    assert "track" not in body

    history = client.get(f"/api/competitions/{cid}/submissions/me").json()["submissions"][0]
    assert history["metrics"]
    assert "result_visibility" not in history
    assert "track" not in history

    board = client.get(f"/api/competitions/{cid}/leaderboard")
    assert board.status_code == 200
    assert "track" not in board.json()

    card = next(
        c for c in client.get("/api/competitions").json()["competitions"] if c["id"] == cid
    )
    assert "my_stats" in card
    assert "my_stats_by_track" not in card
    assert card["tracks"] is None

    login(client)
    export = client.get(f"/api/admin/competitions/{cid}/export.xlsx")
    assert export.status_code == 200, export.text
    assert "-results-" in export.headers["content-disposition"]
    assert "provisional" not in export.headers["content-disposition"]


# --- AI review theo nhánh --------------------------------------------------------------------------


def _snapshot_resources(client, submission: dict) -> list[dict]:
    """Tài nguyên đã chụp trong revision mà bài nộp trỏ tới - đúng ngữ cảnh lượt AI đối chiếu."""

    async def load():
        db = client.app.state.mongo.db
        revision = await db[content_snapshot.REVISIONS_COLLECTION].find_one(
            {"_id": submission["content_snapshot"]["revision_id"]}
        )
        return revision["resources"]

    return asyncio.run(load())


def _dual_ai_competition(
    client, slug: str, *, shared: list[dict], private_resources: list[dict]
) -> dict:
    """Cuộc thi dual đã có thể lệ Markdown, tài nguyên hai tầng và AI bật; trả payload tạo."""
    payload = create_dual_ok(
        client,
        dual_body(slug=slug, private_resources=private_resources) | {"resources": shared},
    )
    _add_rules(client, payload["id"])
    assert enable_ai(client, payload["id"]).status_code == 200
    return payload


def test_bai_v1_dual_chup_ai_snapshot_theo_tai_nguyen_cua_dung_nhanh(client, ai_env):
    """Bài Private đối chiếu với tài nguyên Private, bài Public với danh sách chung.

    Đường chấm ngay trong request (v1) chụp snapshot ở chính lượt nộp. Hai ngữ cảnh tài nguyên
    khác nhau phải cho hai revision khác nhau, nếu không cache AI sẽ phục vụ lại kết luận của
    nhánh này cho nhánh kia.
    """
    shared = {"label": "Chung", "url": "https://drive.google.com/file/d/shared/view"}
    private_only = {"label": "Riêng Private", "url": "https://drive.google.com/file/d/private/view"}
    login(client)
    payload = _dual_ai_competition(
        client, "ai-dual-v1", shared=[shared], private_resources=[private_only]
    )
    publish_dual_v1(client, payload["id"], slug="ai-dual-v1")
    cid = payload["id"]

    login_participant(client)
    private_posted = submit(client, cid, V1_SUBMISSION, track="private")
    assert private_posted.status_code == 201, private_posted.text
    public_posted = submit(client, cid, V1_SUBMISSION, track="public")
    assert public_posted.status_code == 201, public_posted.text

    private = submission_documents(client, {"track": "private"})[0]
    public = submission_documents(client, {"track": "public"})[0]
    assert _snapshot_resources(client, private) == [shared, private_only]
    assert _snapshot_resources(client, public) == [shared]
    assert (
        private["content_snapshot"]["content_hash"] != public["content_snapshot"]["content_hash"]
    )


def test_bai_v2_dual_chup_ai_snapshot_theo_tai_nguyen_cua_dung_nhanh(
    client, fake_runner, ai_env
):
    """Đường queue (v2) ghi bài sau khi chấm xong: snapshot AI vẫn theo đúng nhánh của lượt."""
    shared = {"label": "Chung", "url": "https://drive.google.com/file/d/shared/view"}
    private_only = {"label": "Riêng Private", "url": "https://drive.google.com/file/d/private/view"}
    login(client)
    payload = _dual_ai_competition(
        client, "ai-dual-v2", shared=[shared], private_resources=[private_only]
    )
    publish_dual_v2(client, payload["id"], slug="ai-dual-v2")

    posted = submit(client, payload["id"], V2_SUBMISSION, track="private")
    assert posted.status_code == 202, posted.text
    assert run_worker(client) == 1

    private = submission_documents(client, {"track": "private"})[0]
    assert _snapshot_resources(client, private) == [shared, private_only]
