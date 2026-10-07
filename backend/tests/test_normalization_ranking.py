"""Xếp hạng theo điểm chuẩn hóa: normalize toàn bộ bài hợp lệ trước, chọn đại diện sau.

Bài nộp được ghi thẳng vào Mongo để ghim chính xác thời điểm nhận bài; mỗi lần ghi thẳng như vậy
test phải tự bỏ cache BXH (`_touch`) vì đường API nộp bài thật đã làm việc đó. Luồng chấm thật nằm
ở `test_submissions.py` và `test_scoring_v2_api.py`.
"""

import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from bson import ObjectId
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.competitions.service import COMPETITIONS_COLLECTION
from app.core.datetimes import iso_z
from app.leaderboard import service as leaderboard_service
from app.submissions.service import SUBMISSIONS_COLLECTION
from tests.helpers import (
    login,
    login_participant,
    publish_v2_competition,
    ready_competition,
)

BASE = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)

# Metric nguồn bị ẩn khỏi thí sinh: norm là dữ liệu dẫn xuất nên cũng phải biến mất với họ.
HIDDEN_PRIMARY_NORM_CONTRACT = {
    "metrics": [
        {"key": "accuracy", "label": "Accuracy", "decimals": 4},
        {"key": "n_items", "label": "Số mẫu", "decimals": 0},
    ],
    "primary_metric": "accuracy",
    "higher_is_better": True,
    "visible_metrics": ["n_items"],
}

# Chỉ số nhỏ hơn là tốt hơn: không được hard-code phép so sánh của chiều cao hơn là tốt hơn.
LOSS_CONTRACT = {
    "metrics": [{"key": "loss", "label": "Loss", "decimals": 4}],
    "primary_metric": "loss",
    "higher_is_better": False,
    "visible_metrics": None,
}


def _run(awaitable):
    return asyncio.run(awaitable)


def _participant(client) -> dict:
    return _run(client.app.state.mongo.db["accounts"].find_one({"email": "thi.sinh@vku.vn"}))


def _create_account(client, name: str, email: str) -> ObjectId:
    account_id = ObjectId()
    _run(
        client.app.state.mongo.db["accounts"].insert_one(
            {
                "_id": account_id,
                "email": email,
                "name": name,
                "password_hash": "unused-in-ranking-tests",
                "role": "participant",
                "active": True,
                "created_at": BASE,
                "updated_at": BASE,
            }
        )
    )
    return account_id


def _ready(client, slug: str, *, baseline: float = 0.6) -> dict:
    """Cuộc thi v1 đã publish với chuẩn hóa bật; participant có sẵn đã join."""
    return ready_competition(
        client, slug=slug, normalization={"enabled": True, "baseline": baseline}
    )


def _ready_v2(client, slug: str, *, baseline: float, output_contract: dict) -> dict:
    """Cuộc thi v2 đã publish với chuẩn hóa bật; chạy đủ cổng xác minh bộ chấm như luồng thật."""
    return publish_v2_competition(
        client,
        slug=slug,
        output_contract=output_contract,
        normalization={"enabled": True, "baseline": baseline},
    )


def _submission(
    client,
    competition: dict,
    account_id: ObjectId,
    score: float,
    created_at: datetime,
    *,
    metrics: dict | None = None,
) -> ObjectId:
    """Bài `completed` ghi thẳng vào Mongo - không chạy evaluator, chỉ dựng dữ liệu xếp hạng."""
    submission_id = ObjectId()
    _run(
        client.app.state.mongo.db[SUBMISSIONS_COLLECTION].insert_one(
            {
                "_id": submission_id,
                "competition_id": ObjectId(competition["id"]),
                "account_id": account_id,
                "status": "completed",
                "created_at": created_at,
                "metrics": metrics if metrics is not None else {"f1": score},
                "primary_score": score,
            }
        )
    )
    return submission_id


def _touch(competition: dict) -> None:
    """Bỏ cache BXH sau khi test ghi thẳng Mongo; đường API thật tự làm việc này khi bài đổi."""
    leaderboard_service.invalidate_competition(ObjectId(competition["id"]))


def _review(client, submission_id: ObjectId, status: str) -> None:
    login(client)
    body = {"status": status}
    if status == "rejected":
        body["note"] = "Lý do kiểm thử"
    response = client.patch(f"/api/admin/submissions/{submission_id}/review", json=body)
    assert response.status_code == 200, response.text


def _board(client, competition: dict) -> dict:
    login(client)
    response = client.get(f"/api/admin/competitions/{competition['id']}/leaderboard")
    assert response.status_code == 200, response.text
    return response.json()


def _participant_board(client, slug: str, **params) -> dict:
    login_participant(client)
    response = client.get(f"/api/competitions/{slug}/leaderboard", params=params)
    assert response.status_code == 200, response.text
    return response.json()


def _norms(board: dict) -> dict[str, float]:
    return {entry["account_id"]: entry["normalized_score"] for entry in board["entries"]}


def _list_item(client, slug: str) -> dict:
    response = client.get("/api/competitions")
    assert response.status_code == 200
    return {item["slug"]: item for item in response.json()["competitions"]}[slug]


def test_zero_group_is_ordered_by_admission_and_keeps_the_representative_package(client):
    """Cả nhóm dưới baseline: đội nộp sớm hơn đứng trước, đại diện là bài 0 sớm nhất của đội."""
    competition = _ready(client, "norm-zero", baseline=0.6)
    participant = _participant(client)
    rival_id = _create_account(client, "Đội Nhì", "norm-zero-rival@vku.vn")

    early = _submission(client, competition, participant["_id"], 0.20, BASE)
    _submission(client, competition, participant["_id"], 0.59, BASE + timedelta(minutes=10))
    _submission(client, competition, rival_id, 0.50, BASE + timedelta(minutes=5))
    _touch(competition)

    board = _board(client, competition)
    first, second = board["entries"]

    assert [entry["display_name"] for entry in board["entries"]] == ["Thí Sinh", "Đội Nhì"]
    assert [entry["rank"] for entry in board["entries"]] == [1, 2]
    assert first["normalized_score"] == 0.0
    assert second["normalized_score"] == 0.0

    # Không có tiêu chí phụ theo raw: raw, giờ, metrics và id cùng thuộc bài 0 nộp sớm nhất.
    assert first["primary_score"] == 0.20
    assert first["best_submission_id"] == str(early)
    assert first["best_submission_at"] == iso_z(BASE)
    assert first["metrics"] == {"f1": 0.20}
    assert first["total_submissions"] == 2

    # Mẫu số là best thực của cuộc thi, không suy từ các bài đang hiển thị (0.59 > mọi raw trên bảng).
    assert board["normalization"]["reference_best"] == 0.59
    assert max(entry["primary_score"] for entry in board["entries"]) == 0.50


def test_positive_norm_beats_zero_and_switches_the_representative(client):
    """Đội có bài vượt baseline lấy đại diện mới, dù trước đó đang đứng nhờ bài 0 sớm nhất."""
    competition = _ready(client, "norm-switch", baseline=0.6)
    participant = _participant(client)
    rival_id = _create_account(client, "Đội Dương", "norm-switch-rival@vku.vn")

    _submission(client, competition, participant["_id"], 0.20, BASE)
    _submission(client, competition, rival_id, 0.65, BASE + timedelta(minutes=5))
    later = _submission(
        client, competition, participant["_id"], 0.90, BASE + timedelta(minutes=20)
    )
    _touch(competition)

    board = _board(client, competition)
    first, second = board["entries"]

    assert first["account_id"] == str(participant["_id"])
    assert first["normalized_score"] == pytest.approx(100.0)
    assert first["primary_score"] == 0.90
    assert first["best_submission_id"] == str(later)
    assert first["best_submission_at"] == iso_z(BASE + timedelta(minutes=20))
    assert second["normalized_score"] == pytest.approx(100 * (0.65 - 0.6) / (0.9 - 0.6))
    assert board["normalization"]["reference_best"] == 0.90


def test_new_global_best_recomputes_the_board_and_reject_rolls_it_back(client):
    """Best mới và từ chối bài best đều chạy lại cả bảng, không chỉ đụng đội vừa đổi."""
    competition = _ready(client, "norm-best", baseline=0.5)
    participant = _participant(client)
    mid_id = _create_account(client, "Đội Tốt", "norm-best-mid@vku.vn")
    top_id = _create_account(client, "Đội Đỉnh", "norm-best-top@vku.vn")

    _submission(client, competition, participant["_id"], 0.50, BASE)
    _submission(client, competition, mid_id, 0.80, BASE + timedelta(minutes=1))
    top = _submission(client, competition, top_id, 1.00, BASE + timedelta(minutes=2))
    _touch(competition)

    board = _board(client, competition)
    norms = _norms(board)
    assert board["entries"][0]["account_id"] == str(top_id)
    assert norms[str(top_id)] == pytest.approx(100.0)
    assert norms[str(mid_id)] == pytest.approx(60.0)  # 100 × (0,80 − 0,50) ÷ (1,00 − 0,50)
    assert norms[str(participant["_id"])] == 0.0  # bằng baseline cũng là 0
    assert board["normalization"]["reference_best"] == 1.00

    # Từ chối bài best duy nhất: mẫu số lùi về 0.80 và mọi norm được tính lại.
    _review(client, top, "rejected")
    board = _board(client, competition)
    norms = _norms(board)
    assert str(top_id) not in norms
    assert board["entries"][0]["account_id"] == str(mid_id)
    assert norms[str(mid_id)] == pytest.approx(100.0)
    assert norms[str(participant["_id"])] == 0.0
    assert board["normalization"]["reference_best"] == 0.80

    # Khôi phục: bài về bảng với timestamp gốc và lấy lại đúng vị trí cũ.
    _review(client, top, "accepted")
    board = _board(client, competition)
    assert [entry["account_id"] for entry in board["entries"]] == [
        str(top_id),
        str(mid_id),
        str(participant["_id"]),
    ]
    assert board["entries"][0]["best_submission_id"] == str(top)
    assert board["normalization"]["reference_best"] == 1.00
    assert _norms(board)[str(mid_id)] == pytest.approx(60.0)


def test_rejecting_the_earliest_zero_falls_back_and_restore_keeps_its_time(client):
    """Bài 0 sớm nhất rời bảng thì đội lùi về bài kế tiếp; khôi phục lấy lại timestamp gốc."""
    competition = _ready(client, "norm-fallback", baseline=0.6)
    participant = _participant(client)
    rival_id = _create_account(client, "Đội Giữa", "norm-fallback-rival@vku.vn")

    first = _submission(client, competition, participant["_id"], 0.20, BASE)
    second = _submission(
        client, competition, participant["_id"], 0.40, BASE + timedelta(minutes=30)
    )
    _submission(client, competition, rival_id, 0.30, BASE + timedelta(minutes=5))
    _touch(competition)

    board = _board(client, competition)
    assert [entry["account_id"] for entry in board["entries"]] == [
        str(participant["_id"]),
        str(rival_id),
    ]
    assert board["entries"][0]["best_submission_id"] == str(first)

    _review(client, first, "rejected")
    board = _board(client, competition)
    assert [entry["account_id"] for entry in board["entries"]] == [
        str(rival_id),
        str(participant["_id"]),
    ]
    assert board["entries"][1]["best_submission_id"] == str(second)
    assert board["entries"][1]["best_submission_at"] == iso_z(BASE + timedelta(minutes=30))

    _review(client, first, "accepted")
    board = _board(client, competition)
    assert [entry["account_id"] for entry in board["entries"]] == [
        str(participant["_id"]),
        str(rival_id),
    ]
    assert board["entries"][0]["best_submission_id"] == str(first)
    assert board["entries"][0]["best_submission_at"] == iso_z(BASE)


def test_equal_norm_and_equal_admission_tie_breaks_deterministically(client):
    """Hai đội cùng norm, cùng giờ nhận bài: thứ tự theo account_id và không đổi giữa các lượt đọc."""
    competition = _ready(client, "norm-tie", baseline=0.6)
    participant = _participant(client)
    rival_id = _create_account(client, "Đội Trùng Giờ", "norm-tie-rival@vku.vn")

    _submission(client, competition, participant["_id"], 0.55, BASE)
    _submission(client, competition, rival_id, 0.58, BASE)
    _touch(competition)

    first = _board(client, competition)
    second = _board(client, competition)
    order = [entry["account_id"] for entry in first["entries"]]
    assert order == sorted(order)
    assert [entry["account_id"] for entry in second["entries"]] == order
    assert [entry["normalized_score"] for entry in first["entries"]] == [0.0, 0.0]


def test_lower_is_better_uses_the_minimum_as_reference(client, fake_runner):
    """Chiều nhỏ-hơn-tốt-hơn: mẫu số là min, cả bảng tính lại khi có min mới."""
    fake_runner.metrics = {"loss": 0.8}
    competition = _ready_v2(client, "norm-loss", baseline=0.9, output_contract=LOSS_CONTRACT)
    participant = _participant(client)
    mid_id = _create_account(client, "Đội Loss", "norm-loss-mid@vku.vn")
    min_id = _create_account(client, "Đội Min", "norm-loss-min@vku.vn")

    _submission(
        client, competition, participant["_id"], 0.9, BASE, metrics={"loss": 0.9}
    )
    _submission(
        client, competition, mid_id, 0.7, BASE + timedelta(minutes=5), metrics={"loss": 0.7}
    )
    _submission(
        client, competition, min_id, 0.5, BASE + timedelta(minutes=10), metrics={"loss": 0.5}
    )
    _touch(competition)

    board = _board(client, competition)
    assert board["normalization"]["higher_is_better"] is False
    assert board["normalization"]["source_metric"] == "loss"
    assert board["normalization"]["reference_best"] == 0.5
    assert board["entries"][0]["account_id"] == str(min_id)
    norms = _norms(board)
    assert norms[str(min_id)] == pytest.approx(100.0)
    assert norms[str(mid_id)] == pytest.approx(50.0)  # 100 × (0,9 − 0,7) ÷ (0,9 − 0,5)
    assert norms[str(participant["_id"])] == 0.0  # bằng baseline là 0, không phải 100

    # Min mới: mẫu số đổi nên đội không nộp gì cũng bị tính lại norm.
    _submission(
        client, competition, participant["_id"], 0.3, BASE + timedelta(minutes=15), metrics={"loss": 0.3}
    )
    _touch(competition)
    board = _board(client, competition)
    assert board["normalization"]["reference_best"] == 0.3
    norms = _norms(board)
    assert norms[str(participant["_id"])] == pytest.approx(100.0)
    assert norms[str(min_id)] == pytest.approx(100 * (0.9 - 0.5) / (0.9 - 0.3))
    assert norms[str(mid_id)] == pytest.approx(100 * (0.9 - 0.7) / (0.9 - 0.3))


def test_pagination_keeps_global_rank_and_norm_for_me(client):
    """Cắt trang chỉ ảnh hưởng hiển thị: hạng và norm của `me` lấy từ toàn bảng."""
    competition = _ready(client, "norm-page", baseline=0.5)
    participant = _participant(client)
    second_id = _create_account(client, "Đội Hai", "norm-page-second@vku.vn")
    third_id = _create_account(client, "Đội Ba", "norm-page-third@vku.vn")

    _submission(client, competition, participant["_id"], 0.50, BASE)
    _submission(client, competition, second_id, 0.75, BASE + timedelta(minutes=5))
    _submission(client, competition, third_id, 1.00, BASE + timedelta(minutes=10))
    _touch(competition)

    page = _participant_board(client, "norm-page", limit=1)
    assert page["total"] == 3
    assert page["has_more"] is True
    assert [entry["rank"] for entry in page["entries"]] == [1]
    assert page["entries"][0]["normalized_score"] == pytest.approx(100.0)
    assert page["me"]["rank"] == 3
    assert page["me"]["normalized_score"] == 0.0
    assert page["me"]["is_current_user"] is True
    assert page["normalization"]["reference_best"] == 1.00


def test_hidden_source_metric_nulls_norm_for_participants_only(client, fake_runner):
    """Metric nguồn bị ẩn: thí sinh không thấy norm lẫn metadata, admin vẫn đủ dữ liệu."""
    fake_runner.metrics = {"accuracy": 0.9, "n_items": 4.0}
    competition = _ready_v2(
        client, "norm-hidden", baseline=0.5, output_contract=HIDDEN_PRIMARY_NORM_CONTRACT
    )
    participant = _participant(client)
    _submission(
        client,
        competition,
        participant["_id"],
        0.9,
        BASE,
        metrics={"accuracy": 0.9, "n_items": 4.0},
    )
    _touch(competition)

    page = _participant_board(client, "norm-hidden")
    entry = page["entries"][0]
    assert entry["primary_score"] is None
    assert entry["normalized_score"] is None
    # Khoá vẫn có để UI biết cuộc thi bật chuẩn hóa, nhưng giá trị thì không.
    assert "normalization" in page
    assert page["normalization"] is None

    board = _board(client, competition)
    assert board["entries"][0]["normalized_score"] == pytest.approx(100.0)
    assert board["normalization"]["source_metric"] == "accuracy"
    assert board["normalization"]["reference_best"] == 0.9


def test_my_stats_reports_current_norm_and_hides_it_with_the_board(client):
    """Thẻ cuộc thi dùng đúng norm hiện tại của BXH và mất norm ngay khi bảng bị ẩn."""
    competition = _ready(client, "norm-stats", baseline=0.5)
    participant = _participant(client)
    rival_id = _create_account(client, "Đội Best", "norm-stats-rival@vku.vn")

    _submission(client, competition, participant["_id"], 0.55, BASE)
    _submission(client, competition, rival_id, 0.70, BASE + timedelta(minutes=5))
    _touch(competition)

    login_participant(client)
    stats = _list_item(client, "norm-stats")["my_stats"]
    assert stats["rank"] == 2
    assert stats["best_score"] == 0.55
    assert stats["best_normalized_score"] == pytest.approx(25.0)  # 100 × (0,55 − 0,50) ÷ (0,70 − 0,50)

    login(client)
    hidden = client.patch(
        f"/api/admin/competitions/{competition['id']}", json={"leaderboard_visible": False}
    )
    assert hidden.status_code == 200, hidden.text
    login_participant(client)
    assert _list_item(client, "norm-stats")["my_stats"]["best_normalized_score"] is None
    assert client.get("/api/competitions/norm-stats/leaderboard").status_code == 403

    # Admin vẫn thấy nguyên bảng và norm dù thí sinh không còn quyền xem.
    board = _board(client, competition)
    assert board["entries"][0]["normalized_score"] == pytest.approx(100.0)
    assert board["normalization"]["baseline"] == 0.5


def test_broken_stored_config_nulls_card_stats_and_fails_the_board_loudly(client):
    """Cấu hình norm hỏng (sửa tay ngoài API): thẻ vẫn mở với số liệu null, BXH thất bại ồn ào."""
    competition = _ready(client, "norm-broken", baseline=0.5)
    participant = _participant(client)
    _submission(client, competition, participant["_id"], 0.70, BASE)
    _touch(competition)
    _run(
        client.app.state.mongo.db[COMPETITIONS_COLLECTION].update_one(
            {"_id": ObjectId(competition["id"])}, {"$set": {"normalization.baseline": None}}
        )
    )

    login_participant(client)
    item = _list_item(client, "norm-broken")
    assert item["normalization"] == {"enabled": True, "baseline": None, "version": 1}
    assert item["my_stats"]["rank"] is None
    assert item["my_stats"]["best_normalized_score"] is None

    with TestClient(client.app, raise_server_exceptions=False) as raw:
        raw.cookies.update(client.cookies)
        board = raw.get("/api/competitions/norm-broken/leaderboard")
    assert board.status_code == 500


def test_corrupt_legacy_score_is_excluded_from_norm_board_and_reference(client):
    """Bài cũ có điểm gốc không đọc được (sửa tay ngoài API) bị loại khỏi bảng norm lẫn mẫu số.

    Không được biến nó thành điểm 0 có hạng, cũng không sửa record gốc.
    """
    competition = _ready(client, "norm-corrupt", baseline=0.5)
    participant = _participant(client)
    rival_id = _create_account(client, "Đội Hỏng", "norm-corrupt-rival@vku.vn")

    _submission(client, competition, participant["_id"], 0.70, BASE)
    broken_id = _submission(client, competition, rival_id, 0.90, BASE)
    _run(
        client.app.state.mongo.db[SUBMISSIONS_COLLECTION].update_one(
            {"_id": broken_id}, {"$set": {"primary_score": None}}
        )
    )
    _touch(competition)

    board = _board(client, competition)
    assert [entry["account_id"] for entry in board["entries"]] == [str(participant["_id"])]
    assert board["normalization"]["reference_best"] == 0.70
    stored = _run(client.app.state.mongo.db[SUBMISSIONS_COLLECTION].find_one({"_id": broken_id}))
    assert stored["primary_score"] is None


def test_rejecting_every_submission_empties_the_board_without_error(client):
    """Bảng rỗng vì bị từ chối hết: không NaN, không mẫu số mồ côi, khôi phục là có lại."""
    competition = _ready(client, "norm-empty", baseline=0.6)
    participant = _participant(client)
    only = _submission(client, competition, participant["_id"], 0.70, BASE)
    _touch(competition)

    page = _participant_board(client, "norm-empty")
    assert page["entries"][0]["normalized_score"] == pytest.approx(100.0)

    _review(client, only, "rejected")
    page = _participant_board(client, "norm-empty")
    assert page["entries"] == []
    assert page["total"] == 0
    assert page["me"] is None
    assert page["normalization"]["reference_best"] is None

    _review(client, only, "accepted")
    page = _participant_board(client, "norm-empty")
    assert page["total"] == 1
    assert page["me"]["rank"] == 1
    assert page["me"]["normalized_score"] == pytest.approx(100.0)


def test_export_adds_current_norm_column_and_info_rows(client):
    """File export giữ nguyên điểm gốc, thêm cột norm hiện tại và mặt bằng ở sheet Info."""
    competition = _ready(client, "norm-export", baseline=0.5)
    participant = _participant(client)
    rival_id = _create_account(client, "Đội Xuất", "norm-export-rival@vku.vn")

    _submission(client, competition, participant["_id"], 0.55, BASE)
    rival_submission = _submission(
        client, competition, rival_id, 0.70, BASE + timedelta(minutes=5)
    )
    _touch(competition)

    login(client)
    response = client.get(f"/api/admin/competitions/{competition['id']}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Results"]
    header = [cell.value for cell in sheet[1]]
    assert header[3:5] == ["Best score", "Norm hiện tại (0–100)"]
    assert header[-2:] == ["Best submission time", "Total submissions"]

    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == str(rival_id)
    assert rows[0][3] == 0.70
    assert rows[0][4] == pytest.approx(100.0)
    # Điểm gốc lưu nguyên giá trị, không bị cắt theo số chữ số hiển thị của norm.
    assert rows[1][1] == str(participant["_id"])
    assert rows[1][3] == 0.55
    assert rows[1][4] == pytest.approx(25.0)
    assert sheet.cell(row=2, column=5).number_format == "0.00"

    info = {
        row[0]: row[1]
        for row in workbook["Info"].iter_rows(values_only=True)
        if row and row[0]
    }
    assert info["Norm baseline"] == 0.5
    assert info["Norm reference best"] == 0.70
    assert info["Norm source"] == "f1 (Higher is better)"

    # Export bỏ cache: sau khi bài đại diện bị từ chối, file tải ngay phải khớp bảng mới
    # (0.55 trở thành mẫu số nên norm của chính nó là 100, không còn 25.0 của bản cũ).
    _review(client, rival_submission, "rejected")
    refreshed = client.get(f"/api/admin/competitions/{competition['id']}/export.xlsx")
    rows = list(
        load_workbook(BytesIO(refreshed.content))["Results"].iter_rows(
            min_row=2, values_only=True
        )
    )
    assert len(rows) == 1
    assert rows[0][3] == 0.55
    assert rows[0][4] == pytest.approx(100.0)


def test_bang_toan_cuc_admin_gan_norm_metadata_ngoai_hop_dong(client):
    """Metadata `competitions[]` mang cấu hình norm ở khóa riêng, không trộn vào hợp đồng evaluator."""
    competition = _ready(client, "norm-meta")
    account_id = _create_account(client, "Đội Meta", "doi-meta@vku.vn")
    _submission(client, competition, account_id, 0.8, BASE)

    login(client)
    body = client.get("/api/admin/submissions").json()
    item = next(entry for entry in body["competitions"] if entry["id"] == competition["id"])
    assert item["normalization"] == {"enabled": True, "baseline": 0.6, "version": 1}
    # Hợp đồng kết quả vẫn là output của evaluator, không bị nhét norm vào.
    assert "normalization" not in item["result_contract"]
