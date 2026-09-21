"""Xét duyệt bài nộp: admin từ chối kèm lý do hoặc khôi phục, không đụng trục chấm điểm."""

import asyncio
from datetime import datetime, timezone
from io import BytesIO

from bson import ObjectId
from openpyxl import load_workbook

from app.submissions import service
from app.submissions.service import SUBMISSIONS_COLLECTION
from tests.helpers import (
    SCORING_CONFIG,
    account_id_by_email,
    login,
    login_participant,
    membership_document,
    ready_competition,
    submission_documents,
    submit,
)

# Ground truth của `ready_competition`: id 1,2 → 1 và id 3,4 → 0.
PERFECT = b"id,prediction\n1,1\n2,1\n3,0\n4,0\n"  # f1 = 1.0
HALF = b"id,prediction\n1,1\n2,0\n3,1\n4,0\n"  # f1 = 0.5
GLOBAL_URL = "/api/admin/submissions"
REASON = "Notebook dùng kiến trúc không được phép."


def _run(awaitable):
    return asyncio.run(awaitable)


def _review(client, submission_id, payload):
    return client.patch(f"{GLOBAL_URL}/{submission_id}/review", json=payload)


def _document(client, submission_id) -> dict:
    return _run(
        client.app.state.mongo.db[SUBMISSIONS_COLLECTION].find_one(
            {"_id": ObjectId(submission_id)}
        )
    )


def _admin_row(client, submission_id) -> dict:
    body = client.get(GLOBAL_URL).json()
    return next(row for row in body["submissions"] if row["id"] == submission_id)


def _history_row(client, competition_id, submission_id) -> dict:
    body = client.get(f"/api/competitions/{competition_id}/submissions/me").json()
    return next(item for item in body["submissions"] if item["id"] == submission_id)


def _submit_pair(client, competition_id) -> tuple[str, str]:
    """Nộp hai bài của cùng một participant; trả (bài điểm cao, bài điểm thấp)."""
    return submit(client, competition_id, PERFECT).json()["id"], submit(
        client, competition_id, HALF
    ).json()["id"]


def test_review_requires_admin(client):
    anonymous = client.patch(
        f"{GLOBAL_URL}/{ObjectId()}/review", json={"status": "rejected", "note": REASON}
    )
    assert anonymous.status_code == 401
    assert anonymous.json()["error"]["code"] == "UNAUTHORIZED"

    competition = ready_competition(client)
    submission_id = submit(client, competition["id"], PERFECT).json()["id"]
    forbidden = _review(client, submission_id, {"status": "rejected", "note": REASON})
    assert forbidden.status_code == 403
    assert forbidden.json()["error"]["code"] == "FORBIDDEN"
    assert "review" not in _document(client, submission_id)


def test_review_rejects_unknown_id_and_non_completed_record(client):
    competition = ready_competition(client)
    submission_id = submit(client, competition["id"], PERFECT).json()["id"]
    login(client)

    missing = _review(client, str(ObjectId()), {"status": "rejected", "note": REASON})
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "NOT_FOUND"

    malformed = _review(
        client, "khong-phai-objectid", {"status": "rejected", "note": REASON}
    )
    assert malformed.status_code == 404

    # Record cũ mang trạng thái chấm điểm thất bại không thể nhận quyết định xét duyệt.
    _run(
        client.app.state.mongo.db[SUBMISSIONS_COLLECTION].update_one(
            {"_id": ObjectId(submission_id)}, {"$set": {"status": "rejected"}}
        )
    )
    legacy = _review(client, submission_id, {"status": "rejected", "note": REASON})
    assert legacy.status_code == 422
    assert legacy.json()["error"]["code"] == "INVALID_TRANSITION"


def test_review_validates_status_and_note(client):
    competition = ready_competition(client)
    submission_id = submit(client, competition["id"], PERFECT).json()["id"]
    login(client)

    for payload in (
        {"status": "completed", "note": REASON},
        {"status": "rejected"},
        {"status": "rejected", "note": "   "},
        {"status": "rejected", "note": "x" * (service.MAX_REVIEW_NOTE_LENGTH + 1)},
        {"status": "accepted", "note": REASON},
    ):
        response = _review(client, submission_id, payload)
        assert response.status_code == 422, payload
        assert response.json()["error"]["code"] == "VALIDATION_ERROR"

    # Không payload nào ở trên được phép ghi lại quyết định duyệt.
    assert "review" not in _document(client, submission_id)


def test_reject_keeps_document_artifacts_and_quota(client, fake_artifact_storage):
    competition = ready_competition(client, quota=5)
    high, _ = _submit_pair(client, competition["id"])
    before = _document(client, high)
    quota_before = membership_document(client, competition["id"])["quota_used"]

    login(client)
    response = _review(
        client, high, {"status": "rejected", "note": f"  {REASON}  "}
    )
    assert response.status_code == 200

    after = _document(client, high)
    assert after["status"] == "completed"
    assert after["metrics"] == before["metrics"]
    assert after["primary_score"] == before["primary_score"]
    assert after["artifacts"] == before["artifacts"]
    assert after["submission_no"] == before["submission_no"]
    assert after["created_at"] == before["created_at"]
    assert after["review"]["status"] == "rejected"
    # Lý do được trim trước khi lưu.
    assert after["review"]["note"] == REASON
    # Lượt nộp vẫn bị tiêu như trước - từ chối là hậu kiểm, không hoàn quota.
    assert membership_document(client, competition["id"])["quota_used"] == quota_before

    login_participant(client)
    participant_downloads = [
        client.get(f"/api/competitions/{competition['id']}/submissions/{high}/{kind}")
        for kind in ("prediction", "notebook")
    ]
    assert [item.status_code for item in participant_downloads] == [200, 200]

    login(client)
    admin_downloads = [
        client.get(f"{GLOBAL_URL}/{high}/{kind}") for kind in ("prediction", "notebook")
    ]
    assert [item.status_code for item in admin_downloads] == [200, 200]


def test_participant_sees_reason_without_reviewer_identity(client):
    competition = ready_competition(client)
    high, low = _submit_pair(client, competition["id"])
    login(client)
    _review(client, high, {"status": "rejected", "note": REASON})

    login_participant(client)
    rejected = _history_row(client, competition["id"], high)
    assert rejected["review"] == {"status": "rejected", "note": REASON}
    assert "reviewed_by" not in str(rejected)
    assert "reviewed_at" not in str(rejected)
    # Bài chưa từng bị xét duyệt không có khoá `review`.
    assert "review" not in _history_row(client, competition["id"], low)
    # Metrics vẫn hiển thị để participant đối chiếu được vì sao bài bị từ chối.
    assert rejected["metrics"] is not None

    login(client)
    row = _admin_row(client, high)
    assert row["review"]["status"] == "rejected"
    assert row["review"]["note"] == REASON
    assert row["review"]["reviewed_by"] == {
        "id": str(account_id_by_email(client, "admin@vku.vn")),
        "name": "Admin Mét",
        "email": "admin@vku.vn",
    }
    assert row["review"]["reviewed_at"].endswith("Z")
    # Bài chưa xét duyệt trả `null` để UI phân biệt "mặc định hợp lệ" với "đã khôi phục".
    assert _admin_row(client, low)["review"] is None


def test_restore_clears_note_and_keeps_newest_decision(client):
    competition = ready_competition(client)
    high, _ = _submit_pair(client, competition["id"])
    login(client)
    _review(client, high, {"status": "rejected", "note": "Lý do cũ."})

    restored = _review(client, high, {"status": "accepted"})
    assert restored.status_code == 200
    assert restored.json()["submission"]["review"]["status"] == "accepted"

    document = _document(client, high)
    assert document["status"] == "completed"
    assert document["review"]["status"] == "accepted"
    # Chỉ giữ trạng thái gần nhất nên lý do cũ bị xóa hẳn.
    assert document["review"]["note"] is None
    assert document["review"]["reviewed_by"] == account_id_by_email(client, "admin@vku.vn")

    login_participant(client)
    assert "review" not in _history_row(client, competition["id"], high)

    login(client)
    assert _admin_row(client, high)["review"]["status"] == "accepted"


def test_leaderboard_and_export_drop_rejected_and_fall_back(client):
    competition = ready_competition(client)
    high, low = _submit_pair(client, competition["id"])
    leaderboard_url = f"/api/competitions/{competition['id']}/leaderboard"
    export_url = f"/api/admin/competitions/{competition['id']}/export.xlsx"

    login_participant(client)
    baseline = client.get(leaderboard_url).json()
    assert baseline["total"] == 1
    assert baseline["entries"][0]["best_submission_id"] == high
    assert baseline["entries"][0]["total_submissions"] == 2

    login(client)
    _review(client, high, {"status": "rejected", "note": REASON})

    login_participant(client)
    after_reject = client.get(leaderboard_url).json()
    # Bài hợp lệ kế tiếp lên làm bài tốt nhất, và chỉ nó được đếm vào tổng số bài.
    assert after_reject["total"] == 1
    assert after_reject["entries"][0]["best_submission_id"] == low
    assert after_reject["entries"][0]["primary_score"] == 0.5
    assert after_reject["entries"][0]["total_submissions"] == 1
    assert after_reject["me"]["best_submission_id"] == low

    login(client)
    workbook = load_workbook(
        BytesIO(client.get(export_url).content), read_only=True, data_only=False
    )
    rows = list(workbook["Results"].iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][3] == 0.5
    assert rows[1][8] == 1

    # Từ chối nốt bài còn lại: đội không còn bài hợp lệ nên rời bảng xếp hạng hoàn toàn.
    _review(client, low, {"status": "rejected", "note": REASON})
    login_participant(client)
    empty = client.get(leaderboard_url).json()
    assert empty["total"] == 0
    assert empty["entries"] == []
    assert empty["me"] is None

    # Khôi phục dùng lại metrics đã lưu, không chấm lại; bài kia vẫn bị từ chối nên không được đếm.
    login(client)
    _review(client, high, {"status": "accepted"})
    workbook = load_workbook(
        BytesIO(client.get(export_url).content), read_only=True, data_only=False
    )
    rows = list(workbook["Results"].iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][3] == 1.0
    assert rows[1][8] == 1


def test_review_filter_and_completed_stat_excludes_rejected(client):
    competition = ready_competition(client)
    high, low = _submit_pair(client, competition["id"])
    login(client)
    _review(client, high, {"status": "rejected", "note": REASON})

    all_rows = client.get(GLOBAL_URL).json()
    # Tổng vẫn đếm cả bài bị từ chối; "đã chấm điểm" chỉ đếm bài được tính kết quả.
    assert all_rows["stats"] == {
        "total": 2,
        "competitions": 1,
        "teams": 1,
        "completed": 1,
    }
    assert {row["id"] for row in all_rows["submissions"]} == {high, low}

    rejected = client.get(GLOBAL_URL, params={"review": "rejected"}).json()
    assert [row["id"] for row in rejected["submissions"]] == [high]
    assert rejected["stats"]["completed"] == 0

    accepted = client.get(GLOBAL_URL, params={"review": "accepted"}).json()
    assert [row["id"] for row in accepted["submissions"]] == [low]

    combined = client.get(
        GLOBAL_URL,
        params={"review": "rejected", "status": "completed", "q": "Thí Sinh"},
    ).json()
    assert [row["id"] for row in combined["submissions"]] == [high]
    assert (
        client.get(GLOBAL_URL, params={"review": "rejected", "status": "failed"}).json()[
            "total"
        ]
        == 0
    )

    scoped = client.get(
        f"/api/admin/competitions/{competition['id']}/submissions",
        params={"review": "rejected"},
    ).json()
    assert [row["id"] for row in scoped["submissions"]] == [high]

    for url in (GLOBAL_URL, f"/api/admin/competitions/{competition['id']}/submissions"):
        invalid = client.get(url, params={"review": "unknown"})
        assert invalid.status_code == 422
        assert invalid.json()["error"]["code"] == "VALIDATION_ERROR"


def test_review_preserves_retention_and_scoring_lock(client):
    competition = ready_competition(client)
    high, _ = _submit_pair(client, competition["id"])
    login(client)
    _review(client, high, {"status": "rejected", "note": REASON})

    account_id = account_id_by_email(client, "thi.sinh@vku.vn")
    blocked = client.delete(
        f"/api/admin/competitions/{competition['id']}/members/{account_id}"
    )
    assert blocked.status_code == 409
    assert blocked.json()["error"]["code"] == "MEMBER_HAS_SUBMISSIONS"

    # Bài vẫn `completed` nên cấu hình chấm điểm vẫn khóa như trước khi bị từ chối.
    locked = client.put(
        f"/api/admin/competitions/{competition['id']}/scoring", json=SCORING_CONFIG
    )
    assert locked.status_code == 422
    assert locked.json()["error"]["code"] == "SCORING_LOCKED"

    # Từ chối không xoá gì: cả hai document còn nguyên.
    assert len(submission_documents(client)) == 2


def test_concurrent_reviews_never_mix_metadata(client):
    """Ghi trọn object `review` nên note và người duyệt luôn thuộc về cùng một lần thao tác."""
    competition = ready_competition(client)
    high, _ = _submit_pair(client, competition["id"])
    db = client.app.state.mongo.db
    first_admin, second_admin = ObjectId(), ObjectId()

    async def race():
        await asyncio.gather(
            service.set_submission_review(
                db,
                ObjectId(high),
                status="rejected",
                note="A",
                reviewed_by=first_admin,
                now=datetime.now(timezone.utc),
            ),
            service.set_submission_review(
                db,
                ObjectId(high),
                status="rejected",
                note="B",
                reviewed_by=second_admin,
                now=datetime.now(timezone.utc),
            ),
        )

    _run(race())

    review = _document(client, high)["review"]
    assert set(review) == {"status", "note", "reviewed_by", "reviewed_at"}
    assert (review["note"], review["reviewed_by"]) in {
        ("A", first_admin),
        ("B", second_admin),
    }
