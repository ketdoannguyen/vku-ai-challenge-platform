"""Cache ngắn hạn của bảng xếp hạng: TTL, trần kích thước, single-flight, phân quyền và export.

Cache chỉ giữ danh sách xếp hạng thô; các bài test ở đây kiểm cả việc lọc metric theo từng request
lẫn việc export luôn đọc tươi, và việc cache không rò rỉ giữa các DB/test với nhau.
"""

import asyncio
import time
from datetime import datetime, timedelta, timezone
from io import BytesIO

from bson import ObjectId
from mongomock_motor import AsyncMongoMockClient
from openpyxl import load_workbook

from app.competitions.service import COMPETITIONS_COLLECTION
from app.leaderboard import service
from app.submissions.service import SUBMISSIONS_COLLECTION
from tests.helpers import V2_SCHEMA, login, login_participant, publish_v2_competition

BASE = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)

# Admin chỉ cho thí sinh thấy accuracy; f1 là chỉ số chính bị ẩn khỏi mọi payload thí sinh.
HIDDEN_CONTRACT = {
    "metrics": [
        {"key": "accuracy", "label": "Accuracy", "decimals": 4},
        {"key": "f1", "label": "F1", "decimals": 4},
    ],
    "primary_metric": "f1",
    "higher_is_better": True,
    "visible_metrics": ["accuracy"],
}
HIDDEN_METRICS = {"accuracy": 0.9, "f1": 0.8}


def _run(awaitable):
    return asyncio.run(awaitable)


def _account(client, email: str) -> dict:
    return _run(client.app.state.mongo.db["accounts"].find_one({"email": email}))


def _create_account(client, name: str, email: str) -> ObjectId:
    account_id = ObjectId()
    _run(
        client.app.state.mongo.db["accounts"].insert_one(
            {
                "_id": account_id,
                "email": email,
                "name": name,
                "password_hash": "unused-in-cache-tests",
                "role": "participant",
                "active": True,
                "created_at": BASE,
                "updated_at": BASE,
            }
        )
    )
    return account_id


def _competition(client, slug: str, *, leaderboard_visible: bool = True) -> ObjectId:
    competition_id = ObjectId()
    now = datetime.now(timezone.utc)
    _run(
        client.app.state.mongo.db[COMPETITIONS_COLLECTION].insert_one(
            {
                "_id": competition_id,
                "slug": slug,
                "name": slug.replace("-", " ").title(),
                "status": "published",
                "primary_metric": "f1",
                "leaderboard_visible": leaderboard_visible,
                "start_at": now - timedelta(days=1),
                "end_at": now + timedelta(days=1),
            }
        )
    )
    return competition_id


def _submission(
    client, competition_id, account_id, score: float, created_at: datetime
) -> ObjectId:
    submission_id = ObjectId()
    _run(
        client.app.state.mongo.db[SUBMISSIONS_COLLECTION].insert_one(
            {
                "_id": submission_id,
                "competition_id": competition_id,
                "account_id": account_id,
                "status": "completed",
                "metrics": {"f1": score, "precision": score, "recall": score},
                "primary_score": score,
                "created_at": created_at,
            }
        )
    )
    return submission_id


def _v2_competition(competition_id: ObjectId, *, higher_is_better: bool) -> dict:
    """Document v2 tối thiểu để `contracts` đọc ra chiều xếp hạng - không cần chạy qua API."""
    return {
        "_id": competition_id,
        "scoring_config": {
            "version": 2,
            "revision": 1,
            "input_schema": V2_SCHEMA,
            "evaluator": {"name": "Bộ chấm thử"},
            "output_contract": {
                "metrics": [{"key": "loss", "label": "Loss", "decimals": 3}],
                "primary_metric": "loss",
                "higher_is_better": higher_is_better,
            },
        },
    }


def _counting_ranked_entries(monkeypatch) -> dict:
    """Đếm số lần loader thật chạy để biết lượt đọc nào đi qua cache."""
    original = service.ranked_entries
    counter = {"count": 0}

    async def counted(db, competition):
        counter["count"] += 1
        return await original(db, competition)

    monkeypatch.setattr(service, "ranked_entries", counted)
    return counter


async def test_cache_serves_within_ttl_and_reloads_after_expiry(mock_db):
    cache = service._RankedEntriesCache(max_entries=8, ttl_seconds=0.05)
    competition = {"_id": ObjectId(), "primary_metric": "f1"}
    calls = []

    async def loader(db, competition):
        calls.append(len(calls) + 1)
        return [{"load": len(calls)}]

    first = await cache.get(mock_db, competition, loader)
    second = await cache.get(mock_db, competition, loader)
    assert calls == [1]
    assert second == first

    await asyncio.sleep(0.06)
    third = await cache.get(mock_db, competition, loader)
    assert calls == [1, 2]
    assert third == [{"load": 2}]


async def test_cache_is_bounded_and_evicts_the_oldest_entry(mock_db):
    cache = service._RankedEntriesCache(max_entries=2, ttl_seconds=60.0)
    loaded = []

    async def loader(db, competition):
        loaded.append(competition["_id"])
        return [{"competition": competition["_id"]}]

    first, second, third = ({"_id": ObjectId(), "primary_metric": "f1"} for _ in range(3))
    await cache.get(mock_db, first, loader)
    await cache.get(mock_db, second, loader)
    await cache.get(mock_db, third, loader)
    assert loaded == [first["_id"], second["_id"], third["_id"]]

    # Entry cũ nhất bị đẩy khỏi cache nên phải nạp lại; entry mới nhất vẫn còn nguyên.
    await cache.get(mock_db, first, loader)
    await cache.get(mock_db, third, loader)
    assert loaded == [first["_id"], second["_id"], third["_id"], first["_id"]]


async def test_concurrent_misses_share_a_single_load(mock_db):
    cache = service._RankedEntriesCache(max_entries=8, ttl_seconds=60.0)
    competition = {"_id": ObjectId(), "primary_metric": "f1"}
    started = asyncio.Event()
    release = asyncio.Event()
    calls = []

    async def loader(db, competition):
        calls.append(len(calls) + 1)
        started.set()
        await release.wait()
        return [{"rank": 1}]

    tasks = [
        asyncio.create_task(cache.get(mock_db, competition, loader)) for _ in range(8)
    ]
    await started.wait()
    # Nhường lượt để 7 task còn lại thấy future đang bay và chờ chung thay vì tự nạp.
    await asyncio.sleep(0)
    assert calls == [1]
    release.set()
    results = await asyncio.gather(*tasks)
    assert calls == [1]
    assert all(result == [{"rank": 1}] for result in results)

    await cache.get(mock_db, competition, loader)
    assert calls == [1]


async def test_failed_load_reaches_waiters_without_poisoning_the_cache(mock_db):
    cache = service._RankedEntriesCache(max_entries=8, ttl_seconds=60.0)
    competition = {"_id": ObjectId(), "primary_metric": "f1"}
    started = asyncio.Event()
    release = asyncio.Event()
    calls = []

    async def loader(db, competition):
        calls.append(len(calls) + 1)
        if len(calls) == 1:
            started.set()
            await release.wait()
            raise RuntimeError("truy vấn hỏng")
        return [{"rank": 1}]

    tasks = [
        asyncio.create_task(cache.get(mock_db, competition, loader)) for _ in range(4)
    ]
    await started.wait()
    await asyncio.sleep(0)
    release.set()
    results = await asyncio.gather(*tasks, return_exceptions=True)
    assert calls == [1]
    assert all(isinstance(result, RuntimeError) for result in results)

    # Lỗi không nằm lại trong cache: lượt đọc sau thử lại từ đầu và thành công.
    assert await cache.get(mock_db, competition, loader) == [{"rank": 1}]
    assert calls == [1, 2]


async def test_invalidation_during_inflight_load_never_serves_or_stores_stale(mock_db):
    """Invalidate trong lúc loader đang bay: người đọc sau phải nạp lại, bản cũ không quay lại cache.

    Kịch bản thật: admin duyệt/từ chối bài trong khi một lượt đọc bảng xếp hạng đang truy vấn. Người
    đọc sau quyết định duyệt không được chờ chung lượt truy vấn cũ, và lượt cũ kết thúc muộn cũng
    không được ghi lại vào cache.
    """
    cache = service._RankedEntriesCache(max_entries=8, ttl_seconds=60.0)
    competition = {"_id": ObjectId(), "primary_metric": "f1"}
    stale_started = asyncio.Event()
    release_stale = asyncio.Event()
    release_fresh = asyncio.Event()
    calls = []

    async def loader(db, competition):
        calls.append(len(calls) + 1)
        if len(calls) == 1:
            stale_started.set()
            await release_stale.wait()
            return [{"load": "stale"}]
        await release_fresh.wait()
        return [{"load": "fresh"}]

    # Lượt nạp cũ bị chặn giữa chừng; một người chờ chung vào hàng trước invalidate.
    stale_reader = asyncio.create_task(cache.get(mock_db, competition, loader))
    await stale_started.wait()
    pre_waiter = asyncio.create_task(cache.get(mock_db, competition, loader))
    for _ in range(3):
        await asyncio.sleep(0)
    assert calls == [1]  # pre_waiter đã vào hàng chờ chung thay vì tự nạp

    cache.invalidate(competition["_id"])

    # Hai người đọc sau invalidate không được chờ chung lượt nạp cũ: phải có lượt nạp mới.
    fresh_reader = asyncio.create_task(cache.get(mock_db, competition, loader))
    fresh_waiter = asyncio.create_task(cache.get(mock_db, competition, loader))
    for _ in range(3):
        await asyncio.sleep(0)
    assert calls == [1, 2]  # lượt nạp mới đã bắt đầu, hai người đọc mới chờ chung nó

    release_fresh.set()
    assert await fresh_reader == [{"load": "fresh"}]
    assert await fresh_waiter == [{"load": "fresh"}]
    release_stale.set()
    # Người đọc bắt đầu trước invalidate vẫn nhận bản của thời điểm họ đọc, không bị treo.
    assert await stale_reader == [{"load": "stale"}]
    assert await pre_waiter == [{"load": "stale"}]

    # Lượt cũ kết thúc sau invalidate không được ghi đè cache: lượt đọc kế tiếp vẫn là bản mới.
    assert await cache.get(mock_db, competition, loader) == [{"load": "fresh"}]
    assert calls == [1, 2]


async def test_cache_keys_separate_databases_competitions_and_rank_direction():
    first_db = AsyncMongoMockClient().test_db
    second_db = AsyncMongoMockClient().test_db
    cache = service._RankedEntriesCache(max_entries=8, ttl_seconds=60.0)
    loaded = []

    async def loader(db, competition):
        loaded.append(competition["_id"])
        return [{"load": len(loaded)}]

    competition_id = ObjectId()
    ascending = _v2_competition(competition_id, higher_is_better=False)
    descending = _v2_competition(competition_id, higher_is_better=True)
    assert service._rank_direction(ascending) == 1
    assert service._rank_direction(descending) == -1

    await cache.get(first_db, ascending, loader)
    await cache.get(second_db, ascending, loader)  # DB khác -> nạp riêng
    await cache.get(first_db, descending, loader)  # cùng DB, khác chiều -> nạp riêng
    assert len(loaded) == 3

    assert await cache.get(first_db, ascending, loader) == [{"load": 1}]
    assert len(loaded) == 3


async def test_invalidation_leaves_other_databases_and_competitions_cached():
    """Invalidate một cuộc thi không đụng cache của cuộc thi khác hay DB khác."""
    first_db = AsyncMongoMockClient().test_db
    second_db = AsyncMongoMockClient().test_db
    cache = service._RankedEntriesCache(max_entries=8, ttl_seconds=60.0)
    loaded = []

    async def loader(db, competition):
        loaded.append((db.name, competition["_id"]))
        return [{"competition": competition["_id"]}]

    target = {"_id": ObjectId(), "primary_metric": "f1"}
    other = {"_id": ObjectId(), "primary_metric": "f1"}
    await cache.get(first_db, target, loader)
    await cache.get(first_db, other, loader)
    await cache.get(second_db, other, loader)
    assert len(loaded) == 3

    cache.invalidate(target["_id"])

    # Cuộc thi khác trên cùng DB và trên DB khác vẫn nguyên cache, không nạp lại.
    assert await cache.get(first_db, other, loader) == [{"competition": other["_id"]}]
    assert await cache.get(second_db, other, loader) == [{"competition": other["_id"]}]
    assert len(loaded) == 3
    # Chỉ cuộc thi bị invalidate mới phải nạp lại.
    assert await cache.get(first_db, target, loader) == [{"competition": target["_id"]}]
    assert len(loaded) == 4


def test_shared_cache_never_leaks_hidden_metrics_to_participant(client, fake_runner, monkeypatch):
    fake_runner.metrics = dict(HIDDEN_METRICS)
    competition = publish_v2_competition(
        client, slug="cache-an-metric", output_contract=HIDDEN_CONTRACT
    )
    competition_id = ObjectId(competition["id"])
    account_id = _create_account(client, "Đội Ẩn", "cache-hidden@vku.vn")
    _run(
        client.app.state.mongo.db[SUBMISSIONS_COLLECTION].insert_one(
            {
                "_id": ObjectId(),
                "competition_id": competition_id,
                "account_id": account_id,
                "status": "completed",
                "metrics": dict(HIDDEN_METRICS),
                "primary_score": 0.8,
                "created_at": BASE,
            }
        )
    )
    counter = _counting_ranked_entries(monkeypatch)

    # Client đang là thí sinh sau publish_v2_competition: chỉ thấy accuracy.
    board = client.get(f"/api/competitions/{competition_id}/leaderboard").json()
    assert counter["count"] == 1
    assert board["primary_metric"] is None
    assert board["entries"][0]["metrics"] == {"accuracy": 0.9}
    assert board["entries"][0]["primary_score"] is None
    assert "account_id" not in board["entries"][0]

    # Admin đọc đúng entry cache đó nhưng nhận dữ liệu thô đầy đủ - cache không mang quyền của ai.
    login(client)
    admin_board = client.get(f"/api/admin/competitions/{competition_id}/leaderboard").json()
    assert counter["count"] == 1
    assert admin_board["primary_metric"] == "f1"
    assert admin_board["entries"][0]["metrics"] == {"accuracy": 0.9, "f1": 0.8}
    assert admin_board["entries"][0]["primary_score"] == 0.8
    assert admin_board["entries"][0]["account_id"] == str(account_id)


def test_export_bypasses_the_cache_and_reads_fresh_entries(client, monkeypatch):
    participant = _account(client, "thi.sinh@vku.vn")
    newcomer_id = _create_account(client, "Đội Mới", "cache-export-new@vku.vn")
    competition_id = _competition(client, "cache-export")
    _submission(client, competition_id, participant["_id"], 0.6, BASE)
    counter = _counting_ranked_entries(monkeypatch)

    login_participant(client)
    board_url = f"/api/competitions/{competition_id}/leaderboard"
    assert client.get(board_url).json()["total"] == 1
    assert counter["count"] == 1

    # Bài mới hoàn thành sau khi bảng đã vào cache: bảng còn phục vụ bản cũ trong TTL...
    _submission(client, competition_id, newcomer_id, 0.9, BASE + timedelta(minutes=1))
    assert client.get(board_url).json()["total"] == 1
    assert counter["count"] == 1

    # ...nhưng file export luôn đọc tươi nên có đủ hai đội.
    login(client)
    export = client.get(f"/api/admin/competitions/{competition_id}/export.xlsx")
    assert export.status_code == 200
    assert counter["count"] == 2
    rows = list(
        load_workbook(BytesIO(export.content), read_only=True)["Results"].iter_rows(
            values_only=True
        )
    )
    assert [row[2] for row in rows[1:]] == ["Đội Mới", "Thí Sinh"]


def test_leaderboard_needs_auth_and_hidden_board_never_reaches_the_cache(client, monkeypatch):
    participant = _account(client, "thi.sinh@vku.vn")
    visible_id = _competition(client, "cache-auth-visible")
    hidden_id = _competition(client, "cache-auth-hidden", leaderboard_visible=False)
    _submission(client, visible_id, participant["_id"], 0.7, BASE)
    _submission(client, hidden_id, participant["_id"], 0.7, BASE)
    counter = _counting_ranked_entries(monkeypatch)

    assert client.get(f"/api/competitions/{visible_id}/leaderboard").status_code == 401
    assert client.get(f"/api/competitions/{hidden_id}/leaderboard").status_code == 401
    assert counter["count"] == 0

    login_participant(client)
    assert client.get(f"/api/competitions/{hidden_id}/leaderboard").status_code == 403
    assert counter["count"] == 0

    login(client)
    assert client.get(f"/api/admin/competitions/{hidden_id}/leaderboard").status_code == 200
    assert counter["count"] == 1


def test_review_decision_invalidates_cached_leaderboard(client):
    participant = _account(client, "thi.sinh@vku.vn")
    competition_id = _competition(client, "cache-review")
    submission_id = _submission(client, competition_id, participant["_id"], 0.8, BASE)
    board_url = f"/api/competitions/{competition_id}/leaderboard"

    login_participant(client)
    assert client.get(board_url).json()["total"] == 1

    # Từ chối: bài rời bảng ngay ở lượt đọc kế tiếp, không chờ hết TTL.
    login(client)
    rejected = client.patch(
        f"/api/admin/submissions/{submission_id}/review",
        json={"status": "rejected", "note": "Không hợp lệ."},
    )
    assert rejected.status_code == 200, rejected.text
    login_participant(client)
    assert client.get(board_url).json()["total"] == 0

    # Khôi phục: bài trở lại bảng cũng ngay lập tức.
    login(client)
    accepted = client.patch(
        f"/api/admin/submissions/{submission_id}/review", json={"status": "accepted"}
    )
    assert accepted.status_code == 200, accepted.text
    login_participant(client)
    assert client.get(board_url).json()["total"] == 1


def test_leaderboard_refreshes_after_ttl(client, monkeypatch):
    competitor_id = _create_account(client, "Đội Sớm", "cache-ttl-early@vku.vn")
    competition_id = _competition(client, "cache-ttl")
    _submission(client, competition_id, competitor_id, 0.5, BASE)

    cache = service._RankedEntriesCache(max_entries=8, ttl_seconds=0.5)
    monkeypatch.setattr(service, "_ranked_entries_cache", cache)

    login_participant(client)
    board_url = f"/api/competitions/{competition_id}/leaderboard"
    assert client.get(board_url).json()["total"] == 1

    _submission(
        client,
        competition_id,
        _account(client, "thi.sinh@vku.vn")["_id"],
        0.9,
        BASE + timedelta(minutes=1),
    )
    assert client.get(board_url).json()["total"] == 1  # còn trong TTL, vẫn bản cũ

    # Hết TTL: lượt đọc kế tiếp nạp lại và thấy bài mới.
    time.sleep(0.55)
    refreshed = client.get(board_url).json()
    assert refreshed["total"] == 2
    assert refreshed["entries"][0]["display_name"] == "Thí Sinh"
