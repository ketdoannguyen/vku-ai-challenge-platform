"""Sprint 06 submission history, leaderboard, admin view and XLSX export."""

import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO

from bson import ObjectId
from openpyxl import load_workbook

from app.competitions.service import COMPETITIONS_COLLECTION
from app.submissions.service import SUBMISSIONS_COLLECTION


def _login(client, email="admin@vku.vn", password="adminmatkhau1"):
    response = client.post("/api/auth/login", json={"identifier": email, "password": password})
    assert response.status_code == 200


def _login_participant(client):
    _login(client, "thi.sinh@vku.vn", "thisinhmatkhau1")


def _run(awaitable):
    return asyncio.run(awaitable)


def _account(client, email: str) -> dict:
    return _run(client.app.state.mongo.db["accounts"].find_one({"email": email}))


def _create_account(client, name: str, email: str) -> ObjectId:
    account_id = ObjectId()
    _run(
        client.app.state.mongo.db["accounts"].insert_one(
            {
                "_id": account_id,
                "email": email,
                "name": name,
                "password_hash": "unused-in-results-tests",
                "role": "participant",
                "active": True,
                "created_at": datetime.now(timezone.utc),
                "updated_at": datetime.now(timezone.utc),
            }
        )
    )
    return account_id


def _create_competition(client, slug: str, *, leaderboard_visible: bool = True) -> ObjectId:
    competition_id = ObjectId()
    now = datetime.now(timezone.utc)
    _run(
        client.app.state.mongo.db[COMPETITIONS_COLLECTION].insert_one(
            {
                "_id": competition_id,
                "slug": slug,
                "name": slug.replace("-", " ").title(),
                "status": "published",
                "primary_metric": "f1",
                "leaderboard_visible": leaderboard_visible,
                "start_at": now - timedelta(days=1),
                "end_at": now + timedelta(days=1),
            }
        )
    )
    return competition_id


def _submission(
    client,
    competition_id: ObjectId,
    account_id: ObjectId,
    score: float,
    created_at: datetime,
    *,
    filename: str = "answers.csv",
    status: str = "completed",
    error_code: str | None = None,
    error_message: str | None = None,
) -> ObjectId:
    submission_id = ObjectId()
    document = {
        "_id": submission_id,
        "competition_id": competition_id,
        "account_id": account_id,
        "file_path": f"submissions/{competition_id}/{account_id}/{submission_id}.csv",
        "original_filename": filename,
        "status": status,
        "metrics": {"f1": score, "precision": score - 0.01, "recall": score - 0.02},
        "primary_score": score,
        "created_at": created_at,
    }
    if error_code:
        document["error_code"] = error_code
    if error_message:
        document["error_message"] = error_message
    _run(client.app.state.mongo.db[SUBMISSIONS_COLLECTION].insert_one(document))
    return submission_id


def test_my_submissions_is_scoped_paginated_newest_first_and_hides_paths(client):
    participant = _account(client, "thi.sinh@vku.vn")
    another_account_id = _create_account(client, "Đội Khác", "other@vku.vn")
    competition_id = _create_competition(client, "history-cup")
    other_competition_id = _create_competition(client, "other-cup")
    base = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)
    first_id = _submission(client, competition_id, participant["_id"], 0.5, base, filename="first.csv")
    latest_id = _submission(
        client,
        competition_id,
        participant["_id"],
        0.6,
        base + timedelta(minutes=2),
        filename="latest.csv",
    )
    _submission(client, competition_id, another_account_id, 0.99, base + timedelta(minutes=3))
    _submission(client, other_competition_id, participant["_id"], 0.98, base + timedelta(minutes=4))

    _login_participant(client)
    response = client.get(
        f"/api/competitions/{competition_id}/submissions/me", params={"limit": 1, "offset": 0}
    )

    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 2
    assert body["limit"] == 1
    assert body["offset"] == 0
    assert [item["id"] for item in body["submissions"]] == [str(latest_id)]
    assert body["submissions"][0]["filename"] == "latest.csv"
    assert body["submissions"][0]["competition_id"] == str(competition_id)
    assert "file_path" not in body["submissions"][0]
    assert "account_id" not in body["submissions"][0]

    second_page = client.get(
        f"/api/competitions/{competition_id}/submissions/me", params={"limit": 1, "offset": 1}
    )
    assert [item["id"] for item in second_page.json()["submissions"]] == [str(first_id)]


def test_my_submissions_returns_safe_reason_for_compatible_failed_record(client):
    participant = _account(client, "thi.sinh@vku.vn")
    competition_id = _create_competition(client, "failed-history")
    failed_id = _submission(
        client,
        competition_id,
        participant["_id"],
        0,
        datetime(2026, 9, 15, tzinfo=timezone.utc),
        status="failed",
        error_code="SCORING_FAILED",
        error_message="Ground truth /data/private/ground_truth.csv không đọc được.",
    )

    _login_participant(client)
    response = client.get(f"/api/competitions/{competition_id}/submissions/me")

    assert response.status_code == 200
    item = response.json()["submissions"][0]
    assert item["id"] == str(failed_id)
    assert item["error"] == {
        "code": "SCORING_FAILED",
        "message": "Không thể chấm điểm bài nộp.",
    }
    assert "file_path" not in item


def test_leaderboard_uses_each_accounts_best_score_and_earlier_best_time(client):
    current = _account(client, "thi.sinh@vku.vn")
    early_account_id = _create_account(client, "Đội Sớm", "early@vku.vn")
    third_account_id = _create_account(client, "Đội Ba", "third@vku.vn")
    competition_id = _create_competition(client, "leaderboard-cup")
    other_competition_id = _create_competition(client, "leaderboard-other")
    base = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)

    _submission(client, competition_id, current["_id"], 0.7, base)
    current_best_id = _submission(client, competition_id, current["_id"], 0.9, base + timedelta(hours=2))
    early_best_id = _submission(client, competition_id, early_account_id, 0.9, base + timedelta(hours=1))
    _submission(client, competition_id, early_account_id, 0.9, base + timedelta(hours=3))
    _submission(client, competition_id, third_account_id, 0.8, base + timedelta(minutes=30))
    _submission(client, other_competition_id, third_account_id, 1.0, base)
    _submission(
        client,
        competition_id,
        early_account_id,
        1.0,
        base,
        status="failed",
    )

    _login_participant(client)
    response = client.get(f"/api/competitions/{competition_id}/leaderboard")

    assert response.status_code == 200
    body = response.json()
    assert body["primary_metric"] == "f1"
    assert body["total"] == 3
    assert [(row["rank"], row["display_name"]) for row in body["entries"]] == [
        (1, "Đội Sớm"),
        (2, "Thí Sinh"),
        (3, "Đội Ba"),
    ]
    assert body["entries"][0]["best_submission_id"] == str(early_best_id)
    assert body["entries"][1]["best_submission_id"] == str(current_best_id)
    assert body["entries"][1]["is_current_user"] is True
    assert sum(row["is_current_user"] for row in body["entries"]) == 1
    assert all("account_id" not in row and "email" not in row for row in body["entries"])


def test_hidden_leaderboard_denies_participant_but_admin_can_view(client):
    participant = _account(client, "thi.sinh@vku.vn")
    competition_id = _create_competition(client, "hidden-cup", leaderboard_visible=False)
    _submission(
        client,
        competition_id,
        participant["_id"],
        0.75,
        datetime(2026, 9, 15, tzinfo=timezone.utc),
    )

    _login_participant(client)
    hidden = client.get(f"/api/competitions/{competition_id}/leaderboard")
    assert hidden.status_code == 403
    assert hidden.json() == {
        "error": {
            "code": "LEADERBOARD_HIDDEN",
            "message": "Bảng xếp hạng hiện chưa được công bố.",
        }
    }

    _login(client)
    admin = client.get(f"/api/admin/competitions/{competition_id}/leaderboard")
    assert admin.status_code == 200
    assert admin.json()["entries"][0]["account_id"] == str(participant["_id"])


def test_admin_submission_list_filters_and_never_exposes_server_path(client):
    participant = _account(client, "thi.sinh@vku.vn")
    another_account_id = _create_account(client, "Đội Khác", "other-filter@vku.vn")
    competition_id = _create_competition(client, "admin-results")
    other_competition_id = _create_competition(client, "admin-results-other")
    base = datetime(2026, 9, 15, tzinfo=timezone.utc)
    wanted_id = _submission(client, competition_id, participant["_id"], 0.8, base)
    _submission(client, competition_id, another_account_id, 0.7, base + timedelta(minutes=1))
    _submission(client, other_competition_id, participant["_id"], 1.0, base + timedelta(minutes=2))

    _login(client)
    response = client.get(
        f"/api/admin/competitions/{competition_id}/submissions",
        params={"q": "Thí Sinh", "status": "completed", "limit": 20, "offset": 0},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 1
    assert body["submissions"][0]["id"] == str(wanted_id)
    assert body["submissions"][0]["account"] == {
        "id": str(participant["_id"]),
        "name": "Thí Sinh",
        "email": "thi.sinh@vku.vn",
    }
    assert "file_path" not in body["submissions"][0]

    invalid_status = client.get(
        f"/api/admin/competitions/{competition_id}/submissions", params={"status": "unknown"}
    )
    assert invalid_status.status_code == 422


def test_admin_xlsx_export_has_rank_values_counts_and_formula_safe_names(client):
    participant = _account(client, "thi.sinh@vku.vn")
    formula_account_id = _create_account(client, "=\x01SUM(1,1)", "formula@vku.vn")
    competition_id = _create_competition(client, "export-cup")
    base = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)
    _submission(client, competition_id, participant["_id"], 0.6, base)
    _submission(client, competition_id, participant["_id"], 0.8, base + timedelta(hours=2))
    _submission(client, competition_id, formula_account_id, 0.9, base + timedelta(hours=1))

    _login(client)
    response = client.get(f"/api/admin/competitions/{competition_id}/export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "export-cup-results-" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
    sheet = workbook["Results"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "Rank",
        "Account ID",
        "Team name",
        "Best score",
        "F1",
        "Precision",
        "Recall",
        "Best submission time",
        "Total submissions",
    )
    assert rows[1][0:4] == (1, str(formula_account_id), "'=SUM(1,1)", 0.9)
    assert rows[1][8] == 1
    assert rows[2][0:4] == (2, str(participant["_id"]), "Thí Sinh", 0.8)
    assert rows[2][8] == 2
    assert all("@vku.vn" not in str(value) for row in rows for value in row if value)
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)


def test_leaderboard_paginates_but_me_stays_global_and_participant_safe(client):
    current = _account(client, "thi.sinh@vku.vn")
    competition_id = _create_competition(client, "paged-cup")
    base = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)
    others = [
        _create_account(client, f"Đội {index}", f"paged{index}@vku.vn") for index in range(1, 5)
    ]
    _submission(client, competition_id, others[0], 0.95, base)
    _submission(client, competition_id, others[1], 0.9, base + timedelta(minutes=1))
    _submission(client, competition_id, others[2], 0.8, base + timedelta(minutes=2))
    _submission(client, competition_id, others[3], 0.7, base + timedelta(minutes=3))
    current_best_id = _submission(
        client, competition_id, current["_id"], 0.6, base + timedelta(minutes=4)
    )
    _submission(client, competition_id, current["_id"], 0.5, base + timedelta(minutes=5))

    _login_participant(client)
    response = client.get(
        f"/api/competitions/{competition_id}/leaderboard", params={"limit": 2, "offset": 2}
    )

    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 5
    assert body["limit"] == 2
    assert body["offset"] == 2
    assert body["has_more"] is True
    # Hạng là thứ hạng toàn cục, không đánh lại số theo trang.
    assert [row["rank"] for row in body["entries"]] == [3, 4]
    assert not any(row["is_current_user"] for row in body["entries"])
    # `me` tìm trên toàn bộ danh sách nên vẫn đúng khi người xem nằm ngoài trang.
    assert body["me"]["rank"] == 5
    assert body["me"]["best_submission_id"] == str(current_best_id)
    assert body["me"]["is_current_user"] is True
    assert all("account_id" not in row and "email" not in row for row in body["entries"])
    assert "account_id" not in body["me"]

    last_page = client.get(
        f"/api/competitions/{competition_id}/leaderboard", params={"limit": 2, "offset": 4}
    ).json()
    assert last_page["has_more"] is False
    assert [row["rank"] for row in last_page["entries"]] == [5]
    assert last_page["entries"][0]["is_current_user"] is True


def test_leaderboard_me_is_null_without_completed_submission(client):
    participant = _account(client, "thi.sinh@vku.vn")
    competition_id = _create_competition(client, "spectator-cup")
    other_account_id = _create_account(client, "Đội Khác", "spectator-other@vku.vn")
    base = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)
    _submission(client, competition_id, other_account_id, 0.9, base)
    # Bài bị từ chối không tính là đã có kết quả nên không được sinh ra `me`.
    _submission(
        client,
        competition_id,
        participant["_id"],
        0.0,
        base + timedelta(minutes=1),
        status="rejected",
        error_code="SUBMISSION_ID_MISMATCH",
    )

    _login_participant(client)
    body = client.get(f"/api/competitions/{competition_id}/leaderboard").json()

    assert body["total"] == 1
    assert body["has_more"] is False
    assert body["me"] is None


def test_leaderboard_rejects_out_of_range_pagination_params(client):
    competition_id = _create_competition(client, "params-cup")
    _login_participant(client)
    url = f"/api/competitions/{competition_id}/leaderboard"

    for params in ({"limit": 0}, {"limit": 201}, {"offset": -1}):
        response = client.get(url, params=params)
        assert response.status_code == 422
        assert response.json()["error"]["code"] == "VALIDATION_ERROR"
