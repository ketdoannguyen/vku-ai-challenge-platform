"""Xếp hạng, Excel và metadata metric đọc theo hợp đồng kết quả của bộ chấm v2.

Submission được ghi thẳng vào Mongo để dựng nhiều đội với điểm khác nhau; luồng chấm thật đã có
`test_scoring_v2_api.py` lo.
"""

import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO

from bson import ObjectId
from openpyxl import load_workbook

from app.competitions.service import COMPETITIONS_COLLECTION
from app.submissions.service import SUBMISSIONS_COLLECTION
from tests.helpers import membership_document, publish_v2_competition

BASE = datetime(2026, 9, 15, 8, tzinfo=timezone.utc)
# Hợp đồng "càng thấp càng tốt": chiều xếp hạng phải đọc từ cấu hình chứ không mặc định giảm dần.
LOSS_CONTRACT = {
    "metrics": [{"key": "loss", "label": "Loss", "decimals": 3}],
    "primary_metric": "loss",
    "higher_is_better": False,
}
# Admin chỉ cho thí sinh thấy accuracy; f1 là chỉ số chính bị ẩn khỏi mọi payload thí sinh.
HIDDEN_PRIMARY_CONTRACT = {
    "metrics": [
        {"key": "accuracy", "label": "Accuracy", "decimals": 4},
        {"key": "f1", "label": "F1", "decimals": 4},
    ],
    "primary_metric": "f1",
    "higher_is_better": True,
    "visible_metrics": ["accuracy"],
}
HIDDEN_PRIMARY_METRICS = {"accuracy": 0.9, "f1": 0.8}


def _run(awaitable):
    return asyncio.run(awaitable)


def _login_admin(client) -> None:
    response = client.post(
        "/api/auth/login", json={"identifier": "admin@vku.vn", "password": "adminmatkhau1"}
    )
    assert response.status_code == 200


def _create_account(client, name: str, email: str) -> ObjectId:
    account_id = ObjectId()
    _run(
        client.app.state.mongo.db["accounts"].insert_one(
            {
                "_id": account_id,
                "email": email,
                "name": name,
                "password_hash": "unused-in-ranking-tests",
                "role": "participant",
                "active": True,
                "created_at": BASE,
                "updated_at": BASE,
            }
        )
    )
    return account_id


def _submission(
    client,
    competition_id,
    account_id,
    *,
    metrics: dict,
    primary_score: float,
    created_at: datetime,
) -> ObjectId:
    submission_id = ObjectId()
    _run(
        client.app.state.mongo.db[SUBMISSIONS_COLLECTION].insert_one(
            {
                "_id": submission_id,
                "competition_id": competition_id,
                "account_id": account_id,
                "status": "completed",
                "metrics": metrics,
                "primary_score": primary_score,
                "created_at": created_at,
            }
        )
    )
    return submission_id


def _loss_cup(client, fake_runner) -> dict:
    """Cuộc thi v2 đã publish với hợp đồng loss và hai đội đã có bài."""
    # Lượt chạy thử lúc publish phải trả đúng khóa của hợp đồng loss.
    fake_runner.metrics = {"loss": 0.2}
    competition = publish_v2_competition(
        client, slug="v2-loss-cup", output_contract=LOSS_CONTRACT
    )
    competition_id = ObjectId(competition["id"])
    better = _create_account(client, "Đội Nhỏ", "doi-nho@vku.vn")
    worse = _create_account(client, "Đội Lớn", "doi-lon@vku.vn")
    _submission(
        client, competition_id, better, metrics={"loss": 0.1}, primary_score=0.1, created_at=BASE
    )
    _submission(
        client,
        competition_id,
        worse,
        metrics={"loss": 0.4},
        primary_score=0.4,
        created_at=BASE + timedelta(minutes=1),
    )
    # Bài thứ hai của đội tốt kém hơn nên không được chọn làm bài tốt nhất.
    _submission(
        client,
        competition_id,
        better,
        metrics={"loss": 0.9},
        primary_score=0.9,
        created_at=BASE + timedelta(minutes=2),
    )
    return competition


def test_lower_is_better_ranks_and_exports_ascending(client, fake_runner):
    competition = _loss_cup(client, fake_runner)

    board = client.get(f"/api/competitions/{competition['id']}/leaderboard")
    assert board.status_code == 200
    assert board.json()["primary_metric"] == "loss"
    assert [
        (entry["display_name"], entry["primary_score"]) for entry in board.json()["entries"]
    ] == [("Đội Nhỏ", 0.1), ("Đội Lớn", 0.4)]

    _login_admin(client)
    export = client.get(f"/api/admin/competitions/{competition['id']}/export.xlsx")
    assert export.status_code == 200
    sheet = load_workbook(BytesIO(export.content), read_only=True)["Results"]
    assert [(row[2], row[3]) for row in sheet.iter_rows(min_row=2, values_only=True)] == [
        ("Đội Nhỏ", 0.1),
        ("Đội Lớn", 0.4),
    ]


def test_export_columns_and_info_sheet_follow_the_contract(client, fake_runner):
    competition = publish_v2_competition(client, slug="v2-export-cup")
    account_id = _create_account(client, "Đội Xuất", "doi-xuat@vku.vn")
    _submission(
        client,
        ObjectId(competition["id"]),
        account_id,
        metrics={"accuracy": 0.75, "n_items": 4.0},
        primary_score=0.75,
        created_at=BASE,
    )

    _login_admin(client)
    export = client.get(f"/api/admin/competitions/{competition['id']}/export.xlsx")
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content))
    results = workbook["Results"]
    rows = list(results.iter_rows(values_only=True))
    # Cột metric lấy từ hợp đồng: nhãn hiển thị của admin, không phải ba cột F1/Precision/Recall.
    assert rows[0] == (
        "Rank",
        "Account ID",
        "Team name",
        "Best score",
        "Accuracy",
        "Số mẫu",
        "Best submission time",
        "Total submissions",
    )
    assert rows[1][4:6] == (0.75, 4.0)
    # Số gốc được giữ nguyên, chỉ định dạng hiển thị theo `decimals`.
    assert results["E2"].number_format == "0.0000"
    assert results["F2"].number_format == "0"

    info = list(workbook["Info"].iter_rows(values_only=True))
    # Ô của dòng một-hai cột được openpyxl đệm None cho đủ bề ngang bảng, nên chỉ so hai cột đầu.
    facts = {(row[0], row[1]) for row in info}
    assert ("Competition", "v2-export-cup") in facts
    assert ("Evaluator", "Bộ chấm thử") in facts
    assert ("Primary metric", "Accuracy (accuracy)") in facts
    assert ("Ranking", "Higher is better") in facts
    assert [row for row in info if row[0] in {"accuracy", "n_items"}] == [
        ("accuracy", "Accuracy", 4),
        ("n_items", "Số mẫu", 0),
    ]
    # Không xuất source, ground truth hay đường dẫn nội bộ.
    assert not any(
        "evaluate" in str(cell) or "ground_truth" in str(cell) for row in info for cell in row
    )


def test_global_admin_list_carries_each_competitions_metric_metadata(client, fake_runner):
    """Bài của hai cuộc thi khác bộ metric vẫn hiển thị đúng nhãn của cuộc thi mình."""
    v2_competition = publish_v2_competition(client, slug="v2-mixed-cup")
    account_id = _create_account(client, "Đội Trộn", "doi-tron@vku.vn")
    _submission(
        client,
        ObjectId(v2_competition["id"]),
        account_id,
        metrics={"accuracy": 0.5, "n_items": 2.0},
        primary_score=0.5,
        created_at=BASE,
    )
    v1_competition_id = ObjectId()
    _run(
        client.app.state.mongo.db[COMPETITIONS_COLLECTION].insert_one(
            {
                "_id": v1_competition_id,
                "slug": "v1-mixed-cup",
                "name": "Cup v1",
                "status": "published",
                "primary_metric": "f1",
                "leaderboard_visible": True,
                "start_at": BASE - timedelta(days=1),
                "end_at": BASE + timedelta(days=1),
            }
        )
    )
    _submission(
        client,
        v1_competition_id,
        account_id,
        metrics={"f1": 0.6, "precision": 0.5, "recall": 0.7},
        primary_score=0.6,
        created_at=BASE + timedelta(minutes=1),
    )

    _login_admin(client)
    body = client.get("/api/admin/submissions").json()
    contracts = {item["id"]: item["result_contract"] for item in body["competitions"]}
    assert [metric["key"] for metric in contracts[v2_competition["id"]]["metrics"]] == [
        "accuracy",
        "n_items",
    ]
    assert contracts[v2_competition["id"]]["primary_metric"] == "accuracy"
    # Cuộc thi v1 giữ đúng ba metric cũ qua adapter, không bị gán hợp đồng của cuộc thi khác.
    assert [metric["key"] for metric in contracts[str(v1_competition_id)]["metrics"]] == [
        "f1",
        "precision",
        "recall",
    ]


def test_whitelist_ap_cho_hop_dong_va_bang_xep_hang_thi_sinh(client, fake_runner):
    """Thí sinh chỉ thấy accuracy: hợp đồng công khai và bảng xếp hạng không lộ f1 đã bị ẩn."""
    # Lượt chạy thử lúc publish phải trả đúng khóa của hợp đồng hai metric.
    fake_runner.metrics = dict(HIDDEN_PRIMARY_METRICS)
    competition = publish_v2_competition(
        client, slug="v2-an-metric", output_contract=HIDDEN_PRIMARY_CONTRACT
    )
    competition_id = ObjectId(competition["id"])
    _submission(
        client,
        competition_id,
        _create_account(client, "Đội Ẩn", "doi-an@vku.vn"),
        metrics=dict(HIDDEN_PRIMARY_METRICS),
        primary_score=0.8,
        created_at=BASE,
    )

    # Client đang là thí sinh sau publish_v2_competition.
    public_config = client.get("/api/competitions/v2-an-metric").json()["submission_config"]
    assert [metric["key"] for metric in public_config["result_contract"]["metrics"]] == ["accuracy"]
    assert public_config["result_contract"]["primary_metric"] is None
    assert public_config["primary_metric"] is None

    board = client.get(f"/api/competitions/{competition_id}/leaderboard").json()
    assert board["primary_metric"] is None
    assert board["entries"][0]["metrics"] == {"accuracy": 0.9}
    assert board["entries"][0]["primary_score"] is None

    # Đường admin vẫn thấy đủ hai metric, chỉ số chính và điểm chính như trước.
    _login_admin(client)
    admin_config = client.get(f"/api/admin/competitions/{competition_id}").json()[
        "submission_config"
    ]
    admin_contract = admin_config["result_contract"]
    assert [metric["key"] for metric in admin_contract["metrics"]] == ["accuracy", "f1"]
    assert admin_contract["primary_metric"] == "f1"
    assert admin_contract["visible_metrics"] == ["accuracy"]
    assert admin_config["primary_metric"] == "f1"
    admin_board = client.get(f"/api/admin/competitions/{competition_id}/leaderboard").json()
    assert admin_board["primary_metric"] == "f1"
    assert admin_board["entries"][0]["metrics"] == {"accuracy": 0.9, "f1": 0.8}
    assert admin_board["entries"][0]["primary_score"] == 0.8


def test_whitelist_ap_cho_lich_su_bai_nop_cua_thi_sinh(client, fake_runner):
    fake_runner.metrics = dict(HIDDEN_PRIMARY_METRICS)
    competition = publish_v2_competition(
        client, slug="v2-an-lich-su", output_contract=HIDDEN_PRIMARY_CONTRACT
    )
    competition_id = ObjectId(competition["id"])
    membership = membership_document(client, competition["id"])
    _submission(
        client,
        competition_id,
        membership["account_id"],
        metrics=dict(HIDDEN_PRIMARY_METRICS),
        primary_score=0.8,
        created_at=BASE,
    )

    history = client.get(f"/api/competitions/{competition_id}/submissions/me").json()
    assert history["submissions"][0]["metrics"] == {"accuracy": 0.9}
    assert history["submissions"][0]["primary_score"] is None
