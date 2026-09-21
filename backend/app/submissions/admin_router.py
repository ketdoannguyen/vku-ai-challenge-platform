"""Admin submission history (per competition và toàn cục), leaderboard, artifact download, XLSX export."""

import logging
import re
from datetime import datetime, timezone
from io import BytesIO

from bson import ObjectId
from bson.errors import InvalidId
from fastapi import APIRouter, Query, Request, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel

from app.accounts.service import ACCOUNTS_COLLECTION
from app.auth.dependencies import AdminAccount
from app.competitions.service import COMPETITIONS_COLLECTION
from app.core.datetimes import iso_z
from app.core.errors import api_error
from app.leaderboard import service as leaderboard_service
from app.submission_artifacts.naming import NOTEBOOK_ARTIFACT, PREDICTION_ARTIFACT
from app.submissions import artifacts as artifacts_reader
from app.submissions import service

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/admin/competitions")
# Bảng submission toàn cục không gắn cuộc thi nào nên tách prefix riêng.
global_router = APIRouter(prefix="/api/admin")
_STATUSES = {"completed", "rejected", "failed"}
_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLSX_ILLEGAL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


class ReviewBody(BaseModel):
    status: str
    note: str | None = None


@router.get("/{competition_id}/submissions")
async def list_submissions(
    competition_id: str,
    request: Request,
    admin: AdminAccount,
    q: str = "",
    status: str | None = None,
    review: str | None = None,
    sort: str = service.DEFAULT_SORT,
    order: str = service.DEFAULT_ORDER,
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
) -> dict:
    db = request.app.state.mongo.db
    competition = await _competition_or_404(db, competition_id)
    _validate_query_params(status, review, sort, order, service.SCOPED_SORT_FIELDS)

    query: dict = {"competition_id": competition["_id"]}
    await _apply_filters(db, query, q, status, review)
    total = await db[service.SUBMISSIONS_COLLECTION].count_documents(query)
    submissions = await service.list_admin_submissions(
        db, query, sort=sort, order=order, limit=limit, offset=offset
    )
    accounts = await _accounts_by_id(db, [item["account_id"] for item in submissions])
    reviewers = await _accounts_by_id(db, _reviewer_ids(submissions))
    return {
        "submissions": [
            _admin_submission(
                item, accounts.get(item["account_id"]), reviewers.get(_reviewer_id(item))
            )
            for item in submissions
        ],
        "total": total,
        "limit": limit,
        "offset": offset,
        "sort": sort,
        "order": order,
    }


@global_router.get("/submissions")
async def list_all_submissions(
    request: Request,
    admin: AdminAccount,
    competition_id: str | None = None,
    q: str = "",
    status: str | None = None,
    review: str | None = None,
    sort: str = service.DEFAULT_SORT,
    order: str = service.DEFAULT_ORDER,
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
) -> dict:
    db = request.app.state.mongo.db
    _validate_query_params(status, review, sort, order, service.SORT_FIELDS)
    query: dict = {}
    if competition_id:
        try:
            query["competition_id"] = ObjectId(competition_id)
        except InvalidId:
            raise api_error(422, "VALIDATION_ERROR", "Cuộc thi không hợp lệ.")
    await _apply_filters(db, query, q, status, review)

    # Bảng toàn cục cần cả bốn số tổng quan nên đếm trong một lượt `$facet`; `total` lấy từ đó.
    stats = await service.submission_stats(db, query)
    total = stats["total"]
    submissions = await service.list_admin_submissions(
        db, query, sort=sort, order=order, limit=limit, offset=offset
    )
    accounts = await _accounts_by_id(db, [item["account_id"] for item in submissions])
    reviewers = await _accounts_by_id(db, _reviewer_ids(submissions))
    competitions = await _competitions_by_id(db, [item["competition_id"] for item in submissions])
    return {
        "submissions": [
            {
                **_admin_submission(
                    item, accounts.get(item["account_id"]), reviewers.get(_reviewer_id(item))
                ),
                "competition": _competition_ref(
                    competitions.get(item["competition_id"]), item["competition_id"]
                ),
            }
            for item in submissions
        ],
        "total": total,
        "limit": limit,
        "offset": offset,
        "sort": sort,
        "order": order,
        "stats": stats,
    }


@global_router.patch("/submissions/{submission_id}/review")
async def set_submission_review(
    submission_id: str, body: ReviewBody, request: Request, admin: AdminAccount
) -> dict:
    """Xét duyệt hậu kiểm một bài đã chấm điểm: từ chối kèm lý do, hoặc khôi phục về hợp lệ.

    Không chấm lại và không hoàn lượt nộp - `status`, metrics, artifact và bộ đếm quota giữ nguyên;
    thao tác chỉ đổi quyết định duyệt. Vì vậy phải tách khỏi `PUT`/`POST` của luồng nộp bài.
    """
    db = request.app.state.mongo.db
    if body.status not in service.REVIEW_STATUSES:
        raise api_error(422, "VALIDATION_ERROR", "Trạng thái duyệt không hợp lệ.")
    note = body.note.strip() if body.note else ""
    if body.status == service.REVIEW_STATUS_REJECTED:
        if not note or len(note) > service.MAX_REVIEW_NOTE_LENGTH:
            raise api_error(
                422,
                "VALIDATION_ERROR",
                f"Lý do không chấp nhận phải có 1-{service.MAX_REVIEW_NOTE_LENGTH} ký tự.",
            )
    elif body.note is not None:
        raise api_error(422, "VALIDATION_ERROR", "Khôi phục bài nộp không nhận lý do.")

    try:
        oid = ObjectId(submission_id)
    except InvalidId:
        raise api_error(404, "NOT_FOUND", "Không tìm thấy bài nộp.")

    updated = await service.set_submission_review(
        db,
        oid,
        status=body.status,
        note=note if body.status == service.REVIEW_STATUS_REJECTED else None,
        reviewed_by=admin["_id"],
        now=datetime.now(timezone.utc),
    )
    if updated is None:
        # Update không khớp vì id sai hoặc vì record chưa chấm được điểm; phân biệt bằng một lượt đọc.
        exists = await db[service.SUBMISSIONS_COLLECTION].count_documents({"_id": oid}, limit=1)
        if not exists:
            raise api_error(404, "NOT_FOUND", "Không tìm thấy bài nộp.")
        raise api_error(
            422, "INVALID_TRANSITION", "Chỉ xét duyệt được bài đã chấm điểm thành công."
        )

    account = await db[ACCOUNTS_COLLECTION].find_one({"_id": updated["account_id"]})
    # Nội dung lý do là dữ liệu của người dùng, không ghi vào log.
    logger.info(
        "Admin %s reviewed submission=%s decision=%s",
        admin["email"],
        oid,
        body.status,
    )
    return {"submission": _admin_submission(updated, account, admin)}


@global_router.get("/submissions/{submission_id}/prediction")
async def admin_download_prediction(
    submission_id: str, request: Request, admin: AdminAccount
) -> Response:
    return await _admin_download(request, submission_id, PREDICTION_ARTIFACT)


@global_router.get("/submissions/{submission_id}/notebook")
async def admin_download_notebook(
    submission_id: str, request: Request, admin: AdminAccount
) -> Response:
    return await _admin_download(request, submission_id, NOTEBOOK_ARTIFACT)


async def _admin_download(request: Request, submission_id: str, kind: str) -> Response:
    db = request.app.state.mongo.db
    try:
        oid = ObjectId(submission_id)
    except InvalidId:
        raise api_error(404, "NOT_FOUND", "Không tìm thấy bài nộp.")
    submission = await db[service.SUBMISSIONS_COLLECTION].find_one({"_id": oid})
    if submission is None:
        raise api_error(404, "NOT_FOUND", "Không tìm thấy bài nộp.")
    competition = await db[COMPETITIONS_COLLECTION].find_one(
        {"_id": submission["competition_id"]}
    )
    if competition is None:
        raise api_error(404, "NOT_FOUND", "Không tìm thấy bài nộp.")
    account = await db[ACCOUNTS_COLLECTION].find_one({"_id": submission["account_id"]})
    return await artifacts_reader.artifact_response(submission, competition, account, kind)


def _validate_query_params(
    status: str | None,
    review: str | None,
    sort: str,
    order: str,
    sort_fields: tuple[str, ...],
) -> None:
    if status is not None and status not in _STATUSES:
        raise api_error(422, "VALIDATION_ERROR", "Trạng thái submission không hợp lệ.")
    # `status` là trạng thái chấm điểm, `review` là quyết định của admin - hai trục, hai tham số.
    if review is not None and review not in service.REVIEW_STATUSES:
        raise api_error(422, "VALIDATION_ERROR", "Trạng thái duyệt không hợp lệ.")
    if sort not in sort_fields:
        raise api_error(422, "VALIDATION_ERROR", "Tiêu chí sắp xếp không hợp lệ.")
    if order not in service.SORT_ORDERS:
        raise api_error(422, "VALIDATION_ERROR", "Thứ tự sắp xếp không hợp lệ.")


async def _apply_filters(
    db, query: dict, q: str, status: str | None, review: str | None
) -> None:
    if status:
        query["status"] = status
    if review:
        query.update(service.review_filter(review))
    account_ids = await service.matching_account_ids(db, q)
    if account_ids is not None:
        query["account_id"] = {"$in": account_ids}


@router.get("/{competition_id}/leaderboard")
async def admin_leaderboard(
    competition_id: str, request: Request, admin: AdminAccount
) -> dict:
    db = request.app.state.mongo.db
    competition = await _competition_or_404(db, competition_id)
    entries = await leaderboard_service.ranked_entries(db, competition["_id"])
    return leaderboard_service.admin_leaderboard_response(competition, entries)


@router.get("/{competition_id}/export.xlsx")
async def export_results(
    competition_id: str, request: Request, admin: AdminAccount
) -> Response:
    db = request.app.state.mongo.db
    competition = await _competition_or_404(db, competition_id)
    entries = await leaderboard_service.ranked_entries(db, competition["_id"])
    content = _build_workbook(entries)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    filename = f"{competition['slug']}-results-{timestamp}.xlsx"
    logger.info(
        "Results exported admin=%s competition=%s entries=%s",
        admin["email"],
        competition["_id"],
        len(entries),
    )
    return Response(
        content,
        media_type=_XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def _competition_or_404(db, competition_id: str) -> dict:
    try:
        object_id = ObjectId(competition_id)
    except InvalidId:
        raise api_error(404, "NOT_FOUND", "Không tìm thấy cuộc thi.")
    competition = await db[COMPETITIONS_COLLECTION].find_one({"_id": object_id})
    if competition is None:
        raise api_error(404, "NOT_FOUND", "Không tìm thấy cuộc thi.")
    return competition


async def _accounts_by_id(db, account_ids: list[ObjectId]) -> dict[ObjectId, dict]:
    return {
        account["_id"]: account
        async for account in db[ACCOUNTS_COLLECTION].find({"_id": {"$in": account_ids}})
    }


async def _competitions_by_id(db, competition_ids: list[ObjectId]) -> dict[ObjectId, dict]:
    return {
        competition["_id"]: competition
        async for competition in db[COMPETITIONS_COLLECTION].find(
            {"_id": {"$in": competition_ids}}
        )
    }


def _admin_submission(
    submission: dict, account: dict | None, reviewer: dict | None = None
) -> dict:
    item = service.submission_history_item(submission)
    item["account"] = {
        "id": str(submission["account_id"]),
        "name": account["name"] if account else "Tài khoản đã xóa",
        "email": account["email"] if account else "",
    }
    # Ghi đè shape participant-safe bằng shape đầy đủ của admin (có người duyệt và thời điểm).
    item[service.REVIEW_FIELD] = _admin_review(submission, reviewer)
    return item


def _reviewer_id(submission: dict):
    """Id admin đã ra quyết định duyệt gần nhất; None nếu bài chưa từng bị xét duyệt."""
    return (submission.get(service.REVIEW_FIELD) or {}).get("reviewed_by")


def _reviewer_ids(submissions: list[dict]) -> list[ObjectId]:
    """Gom theo trang để tra tên người duyệt bằng một truy vấn `$in`, không N+1."""
    return [rid for item in submissions if (rid := _reviewer_id(item)) is not None]


def _admin_review(submission: dict, reviewer: dict | None) -> dict | None:
    """`None` khi bài chưa từng được xét duyệt - UI hiểu là mặc định hợp lệ."""
    review = submission.get(service.REVIEW_FIELD)
    if not review:
        return None
    return {
        "status": review["status"],
        "note": review.get("note"),
        "reviewed_at": iso_z(review["reviewed_at"]),
        "reviewed_by": {
            "id": str(review["reviewed_by"]) if review.get("reviewed_by") else "",
            "name": reviewer["name"] if reviewer else "Tài khoản đã xóa",
            "email": reviewer["email"] if reviewer else "",
        },
    }


def _competition_ref(competition: dict | None, competition_id) -> dict:
    if competition is None:
        return {"id": str(competition_id), "slug": "", "name": "Cuộc thi đã xóa"}
    return {
        "id": str(competition["_id"]),
        "slug": competition["slug"],
        "name": competition["name"],
    }


def _build_workbook(entries: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    headers = (
        "Rank",
        "Account ID",
        "Team name",
        "Best score",
        "F1",
        "Precision",
        "Recall",
        "Best submission time",
        "Total submissions",
    )
    sheet.append(headers)
    for entry in entries:
        metrics = entry["metrics"]
        sheet.append(
            (
                entry["rank"],
                entry["account_id"],
                _formula_safe(entry["display_name"]),
                entry["primary_score"],
                metrics["f1"],
                metrics["precision"],
                metrics["recall"],
                entry["best_submission_at"],
                entry["total_submissions"],
            )
        )

    header_fill = PatternFill("solid", fgColor="DCE8F8")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (8, 26, 28, 14, 12, 12, 12, 24, 18)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _formula_safe(value: str) -> str:
    value = _XLSX_ILLEGAL_CHARACTERS.sub("", value)
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value
